"""#5895 — the file lane reads an owner's Excel form and fills it without losing its look.

The owner brings a quotation form from their own PC; the form is opened only in this
lane (the Chat Worker never imports a workbook parser), and the answer is a *new* file
id. Three things are pinned here on two forms whose labels sit in different places:

* the filled copy carries the mapped values and keeps the form's own styles, merges,
  number formats and sum formulas;
* an item list longer than the form is grown by copying the template row's style;
* the original upload is untouched — its bytes and digest are exactly what was sent.
"""
from __future__ import annotations

import hashlib
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from starlette.testclient import TestClient

from padiem_file_lane.app import create_app

from test_file_lane import SECRET, ticket

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _form(*, label_row: int, label_column: int) -> bytes:
    """One quotation form; the caller decides where its labels sit."""

    book = Workbook()
    sheet = book.active
    sheet.title = "견적서"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    title = sheet.cell(row=1, column=1, value="견 적 서")
    title.font = Font(name="맑은 고딕", size=18, bold=True)
    title.alignment = Alignment(horizontal="center")

    heading = sheet.cell(row=label_row, column=label_column, value="거래처")
    heading.font = Font(bold=True)
    sheet.cell(row=label_row, column=label_column + 1, value="")
    date_cell = sheet.cell(row=label_row + 1, column=label_column, value="날짜")
    date_cell.font = Font(bold=True)

    header_row = label_row + 3
    for column, label in enumerate(("품목", "수량", "단가", "금액"), start=1):
        cell = sheet.cell(row=header_row, column=column, value=label)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.border = Border(bottom=Side(style="thin"))

    item_row = header_row + 1
    for column in range(1, 5):
        sheet.cell(row=item_row, column=column, value="")
    price = sheet.cell(row=item_row, column=3)
    price.number_format = "#,##0"
    for column in range(1, 5):
        sheet.cell(row=item_row, column=column).border = Border(bottom=Side(style="hair"))

    total_row = item_row + 2
    sheet.cell(row=total_row, column=3, value="합계")
    formula = sheet.cell(row=total_row, column=4, value=f"=SUM(D{item_row}:D{total_row - 1})")
    formula.font = Font(bold=True)

    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _upload(client, payload: bytes, name: str = "견적서 양식.xlsx"):
    response = client.post("/files", params={"filename": name}, content=payload,
                           headers=ticket())
    assert response.status_code == 201, response.text
    return response.json()


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(root=tmp_path, secret=SECRET)) as client:
        yield client


@pytest.mark.parametrize("label_row,label_column", [(3, 1), (2, 2)])
def test_the_structure_names_every_non_empty_cell_where_it_sits(client, label_row, label_column):
    payload = _form(label_row=label_row, label_column=label_column)
    meta = _upload(client, payload)
    structure = client.get(f"/files/{meta['file_id']}/structure", headers=ticket()).json()
    assert structure["ok"] is True and structure["name"].endswith(".xlsx")
    sheet = structure["sheets"][0]
    assert sheet["name"] == "견적서" and "A1:D1" in sheet["merged"]
    addresses = {cell["cell"]: cell for cell in sheet["cells"]}
    assert addresses[f"{chr(64 + label_column)}{label_row}"]["text"] == "거래처"
    assert any(cell["formula"] for cell in sheet["cells"])


@pytest.mark.parametrize("label_row,label_column", [(3, 1), (2, 2)])
def test_the_filled_copy_keeps_the_form_and_the_original_is_untouched(client, label_row, label_column):
    payload = _form(label_row=label_row, label_column=label_column)
    meta = _upload(client, payload)
    header_row = label_row + 3
    item_row = header_row + 1
    total_row = item_row + 2
    value_column = chr(64 + label_column + 1)
    mapping = {
        "견적서": {
            f"{value_column}{label_row}": "두암중앙교회",
            f"{value_column}{label_row + 1}": "2026-10-28",
            f"A{item_row}": "LED 100인치",
            f"B{item_row}": 2,
            f"C{item_row}": 1200000,
        }
    }
    response = client.post(f"/files/{meta['file_id']}/fill", headers=ticket(),
                           json={"mapping": mapping, "filename": "견적서_두암중앙교회_2026-10-28.xlsx"})
    assert response.status_code == 201, response.text
    filled_meta = response.json()
    assert filled_meta["file_id"] != meta["file_id"]
    assert filled_meta["name"] == "견적서_두암중앙교회_2026-10-28.xlsx"

    # The owner's original upload is byte-for-byte what was sent.
    original = client.get(f"/files/{meta['file_id']}/raw", headers=ticket()).content
    assert hashlib.sha256(original).hexdigest() == meta["sha256"]

    filled = load_workbook(BytesIO(
        client.get(f"/files/{filled_meta['file_id']}/raw", headers=ticket()).content))
    sheet = filled["견적서"]
    assert sheet.cell(row=label_row, column=label_column + 1).value == "두암중앙교회"
    assert sheet.cell(row=label_row + 1, column=label_column + 1).value == "2026-10-28"
    assert sheet.cell(row=item_row, column=1).value == "LED 100인치"
    assert sheet.cell(row=item_row, column=2).value == 2
    # The form's own look survives: title font, header fill, price format, merge.
    assert "A1:D1" in {str(item) for item in sheet.merged_cells.ranges}
    assert sheet.cell(row=1, column=1).font.bold is True
    assert sheet.cell(row=header_row, column=1).fill.fgColor.rgb == "00DDEBF7"
    assert sheet.cell(row=item_row, column=3).number_format == "#,##0"
    # The form's own sum is left exactly as the owner wrote it.
    assert sheet.cell(row=total_row, column=4).value.startswith("=SUM(")
    filled.close()


def test_a_longer_item_list_grows_with_the_template_row_style(client):
    payload = _form(label_row=3, label_column=1)
    meta = _upload(client, payload)
    # The form has one item row at row 7; the owner needs three.
    response = client.post(f"/files/{meta['file_id']}/fill", headers=ticket(), json={
        "mapping": {"견적서": {"A7": "품목1", "A8": "품목2", "A9": "품목3"}},
        "extend": [{"sheet": "견적서", "template_row": 7, "at_row": 8, "count": 2}],
        "filename": "견적서_늘린항목.xlsx",
    })
    assert response.status_code == 201, response.text
    filled = load_workbook(BytesIO(
        client.get(f"/files/{response.json()['file_id']}/raw", headers=ticket()).content))
    sheet = filled["견적서"]
    assert [sheet.cell(row=row, column=1).value for row in (7, 8, 9)] == ["품목1", "품목2", "품목3"]
    # Each new row wears the template row's own bottom border, so the wider list
    # still reads as the same form rather than as pasted text.
    for row in (8, 9):
        assert sheet.cell(row=row, column=1).border.bottom.style == "hair"
    filled.close()


def test_a_mapping_that_names_nothing_in_the_form_is_refused(client):
    meta = _upload(client, _form(label_row=3, label_column=1))
    for body in (
        {"mapping": {"없는시트": {"A1": "x"}}},
        {"mapping": {"견적서": {"not-a-cell": "x"}}},
        {"mapping": {"견적서": {"A1": {"nested": 1}}}},
        {"mapping": {"견적서": {"A1": "x"}},
         "extend": [{"sheet": "견적서", "template_row": 7, "at_row": 1, "count": 1}]},
        # The form's own sum formula is never overwritten by a mapping.
        {"mapping": {"견적서": {"D9": 999}}},
    ):
        response = client.post(f"/files/{meta['file_id']}/fill", headers=ticket(), json=body)
        assert response.status_code == 422, (body, response.text)


def test_a_file_that_is_not_a_workbook_is_not_a_form(client):
    meta = _upload(client, b"hello", name="hello.txt")
    assert client.get(f"/files/{meta['file_id']}/structure", headers=ticket()).status_code == 415
    assert client.post(f"/files/{meta['file_id']}/fill", headers=ticket(),
                       json={"mapping": {}}).status_code == 415
