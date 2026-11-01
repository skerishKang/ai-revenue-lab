"""Second, fictitious quote layout: real file-lane XLSX fill and PDF comparison.

The mapping below is a controlled model response, not a claim about model accuracy.
The reference PDF is independently created from the same template and facts.
"""
from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from starlette.testclient import TestClient

from padiem_file_lane.app import create_app
from test_file_lane import SECRET, ticket


VALUES = {
    "B6": "Northwind Sample",
    "E6": "2026-11-01",
    "B10": "Display panel",
    "D10": 2,
    "E10": 120000,
    "B11": "Wall mount",
    "D11": 1,
    "E11": 30000,
    "B18": "Delivery within 14 days",
}


def workbook(*, populated=False):
    book = Workbook()
    ws = book.active
    ws.title = "Offer"
    ws.merge_cells("E1:F3")
    ws["E1"] = "LOGO AREA"
    ws["E1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E1"].fill = PatternFill("solid", fgColor="DCEBFA")
    ws.merge_cells("A2:C3")
    ws["A2"] = "SAMPLE HARBOR CO.\nQUOTATION"
    ws["A2"].font = Font(name="Arial", size=15, bold=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    for coordinate, value in {
        "A6": "Customer", "D6": "Date", "A9": "No.", "B9": "Description",
        "D9": "Qty", "E9": "Unit price", "F9": "Amount",
        "E14": "Subtotal", "E15": "VAT 10%", "E16": "Grand total", "A18": "Notes",
    }.items():
        ws[coordinate] = value
        ws[coordinate].font = Font(name="Arial", bold=True)
    for row in (10, 11):
        ws[f"A{row}"] = row - 9
        ws[f"F{row}"] = f"=D{row}*E{row}"
    ws["F14"] = "=SUM(F10:F11)"
    ws["F15"] = "=ROUND(F14*10%,0)"
    ws["F16"] = "=F14+F15"
    for row in ws.iter_rows(min_row=9, max_row=11, min_col=1, max_col=6):
        for cell in row:
            cell.border = Border(bottom=Side(style="thin", color="777777"))
    for row in (10, 11, 14, 15, 16):
        ws[f"F{row}"].number_format = '#,##0'
    for col, width in {"A": 12, "B": 28, "C": 8, "D": 10, "E": 18, "F": 18}.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[18].height = 26
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:F19"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    if populated:
        for cell, value in VALUES.items():
            ws[cell] = value
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def soffice():
    override = os.environ.get("SOFFICE_BIN")
    candidates = [override, shutil.which("soffice"),
                  r"C:\Program Files\LibreOffice\program\soffice.exe"]
    return next((str(p) for p in candidates if p and Path(p).is_file()), None)


def pdf_from_xlsx(path, target):
    exe = soffice()
    if not exe:
        pytest.skip("LibreOffice is required for independent PDF comparison")
    target.mkdir()
    profile = (target / "profile").resolve().as_uri()
    run = subprocess.run(
        [exe, "-env:UserInstallation=" + profile, "--headless", "--convert-to",
         "pdf:calc_pdf_Export", "--outdir", str(target), str(path)],
        capture_output=True, text=True, timeout=90,
    )
    assert run.returncode == 0, run.stderr
    output = target / (path.stem + ".pdf")
    assert output.exists(), run.stdout + run.stderr
    return output


def test_other_company_form_real_file_lane_and_pdf_geometry(tmp_path):
    fitz = pytest.importorskip("fitz")
    source = workbook()
    with TestClient(create_app(root=tmp_path / "lane", secret=SECRET)) as client:
        upload = client.post("/files", params={"filename": "sample-form.xlsx"},
                             content=source, headers=ticket())
        assert upload.status_code == 201
        original_id = upload.json()["file_id"]
        described = client.get(f"/files/{original_id}/structure", headers=ticket())
        assert described.status_code == 200
        structure = described.json()["sheets"][0]
        cells = {cell["cell"]: cell["text"] for cell in structure["cells"]}
        for address, label in {"E1": "LOGO AREA", "A6": "Customer",
                               "B9": "Description", "E14": "Subtotal",
                               "E15": "VAT 10%", "A18": "Notes"}.items():
            assert cells[address] == label
        assert "E1:F3" in structure["merged"]
        result = client.post(f"/files/{original_id}/fill", headers=ticket(), json={
            "mapping": {"Offer": VALUES}, "filename": "filled.xlsx",
        })
        assert result.status_code == 201, result.text
        filled = client.get(f"/files/{result.json()['file_id']}/raw", headers=ticket()).content
        raw = client.get(f"/files/{original_id}/raw", headers=ticket()).content
        assert hashlib.sha256(raw).digest() == hashlib.sha256(source).digest()

    result_book = load_workbook(BytesIO(filled), data_only=False)
    sheet = result_book["Offer"]
    for coordinate, value in VALUES.items():
        assert sheet[coordinate].value == value
    for coordinate, formula in {"F10": "=D10*E10", "F11": "=D11*E11",
                                "F14": "=SUM(F10:F11)", "F15": "=ROUND(F14*10%,0)",
                                "F16": "=F14+F15"}.items():
        assert sheet[coordinate].value == formula
    result_book.close()

    reference_file = tmp_path / "reference.xlsx"
    result_file = tmp_path / "result.xlsx"
    reference_file.write_bytes(workbook(populated=True))
    result_file.write_bytes(filled)
    reference_pdf = pdf_from_xlsx(reference_file, tmp_path / "reference-pdf")
    result_pdf = pdf_from_xlsx(result_file, tmp_path / "result-pdf")
    with fitz.open(reference_pdf) as reference, fitz.open(result_pdf) as output:
        assert len(reference) == len(output) == 1
        assert reference[0].rect == output[0].rect
        assert reference[0].get_text("words") == output[0].get_text("words")
        assert reference[0].get_pixmap().samples == output[0].get_pixmap().samples
        text = output[0].get_text()
        for expected in ("LOGO AREA", "Northwind Sample", "Display panel",
                         "Subtotal", "VAT 10%", "Grand total", "270,000",
                         "27,000", "297,000", "Delivery within 14 days"):
            assert expected in text


def test_file_lane_pdf_matches_reference_and_is_owner_scoped(tmp_path, monkeypatch):
    fitz = pytest.importorskip("fitz")
    from padiem_file_lane import xlsx_pdf
    exe = soffice()
    if not exe:
        pytest.skip("LibreOffice is required for independent PDF comparison")
    monkeypatch.setattr(xlsx_pdf.shutil, "which", lambda _: exe)
    reference_file = tmp_path / "reference.xlsx"
    reference_file.write_bytes(workbook(populated=True))
    reference_pdf = pdf_from_xlsx(reference_file, tmp_path / "reference-output")
    with TestClient(create_app(root=tmp_path / "lane", secret=SECRET)) as client:
        upload = client.post("/files", params={"filename": "form.xlsx"}, content=workbook(), headers=ticket())
        file_id = upload.json()["file_id"]
        filled = client.post(f"/files/{file_id}/fill", headers=ticket(), json={
            "mapping": {"Offer": VALUES}, "filename": "filled.xlsx",
        })
        assert filled.status_code == 201
        response = client.get(f"/files/{filled.json()['file_id']}/pdf", headers=ticket())
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert client.get(f"/files/{filled.json()['file_id']}/pdf").status_code == 401
        with fitz.open(reference_pdf) as expected, fitz.open(stream=response.content, filetype="pdf") as actual:
            assert len(expected) == len(actual) == 1
            assert expected[0].get_text("words") == actual[0].get_text("words")
            assert expected[0].get_pixmap().samples == actual[0].get_pixmap().samples


def test_pdf_unavailable_does_not_prevent_xlsx_download(tmp_path, monkeypatch):
    from padiem_file_lane import xlsx_pdf
    monkeypatch.setattr(xlsx_pdf.shutil, "which", lambda _: None)
    with TestClient(create_app(root=tmp_path / "lane", secret=SECRET)) as client:
        upload = client.post("/files", params={"filename": "form.xlsx"}, content=workbook(), headers=ticket())
        file_id = upload.json()["file_id"]
        response = client.get(f"/files/{file_id}/pdf", headers=ticket())
        assert response.status_code == 503
        assert response.json()["error"]["code"] == "xlsx_pdf_unavailable"
        original = client.get(f"/files/{file_id}/raw", headers=ticket())
        assert original.status_code == 200
        assert original.content.startswith(b"PK")
