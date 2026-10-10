"""#3933 real synthetic XLSX bytes; mock renderer/host, NO Production grants.

The editor and separate verifier use actual openpyxl, not string-only fakes.
They prove semantic cell/style/merge/formula invariants on a synthetic file,
NOT original-byte-identical Excel rendering or live Drive/Claw composition.
"""
from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import unittest

from kagent.artifact_lineage import LineageArtifactRef
from kagent.artifact_registration import register_canonical_artifact
from kagent.contracts import ContractError
from kagent.local_xlsx_artifact_handoff import LocalXlsxArtifactHandoff
from kagent.local_xlsx_revision_pdf import (
    ReviewedCellRevision, LocalRevisedXlsxPdf, MAX_REVISIONS,
    PRODUCTION_DRIVE_WRITE_AUTHORIZED, PRODUCTION_REVISION_HOST_COMPOSED,
    revise_and_render_local_xlsx_pdf,
)
from kagent.xlsx_fidelity_route import classify_xlsx

NOW = datetime(2026, 10, 10, tzinfo=timezone.utc)
WORKSPACE = "workspace_3933_reviewed"
RUN = "run_3933_reviewed"
PDF = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\n%%EOF\n"


def synthetic_source():
    book = Workbook()
    sheet = book.active
    sheet.title = "Quote"
    sheet["A1"] = "Existing School"
    sheet["B2"] = 88_000_000
    sheet["B3"] = "=B2*0.1"
    sheet["B4"] = "=B2+B3"
    sheet["A1"].font = Font(name="Calibri", bold=True)
    sheet.merge_cells("A5:C5")
    sheet.print_area = "A1:C6"
    buf = BytesIO()
    book.save(buf)
    data = buf.getvalue()
    record = register_canonical_artifact(
        artifact_id="quote_source_3933",
        artifact_kind="source.xlsx",
        filename="quote.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=len(data), integrity_ref=sha256(data).hexdigest(),
        workspace_ref=WORKSPACE, run_ref=RUN,
    )
    return LocalXlsxArtifactHandoff(
        record, data, classify_xlsx(data, local_only=True)
    )


EDITS = (
    ReviewedCellRevision("Quote", "A1", "Existing School", "New School"),
    ReviewedCellRevision("Quote", "B2", 88_000_000, 99_200_000),
)


class TrustedEditor:
    def __init__(self):
        self.calls = 0

    def revise_xlsx_copy(self, content, edits):
        self.calls += 1
        wb = load_workbook(BytesIO(content), data_only=False)
        for edit in edits:
            cell = wb[edit.sheet_name][edit.cell]
            assert cell.value == edit.previous_value
            assert cell.data_type != "f"
            cell.value = edit.replacement_value
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()


class IndependentVerifier:
    def __init__(self, accepted=True):
        self.calls = 0
        self.accepted = accepted

    def verify_only_reviewed_edits(self, source, revised, edits):
        self.calls += 1
        if not self.accepted:
            return False
        original = load_workbook(BytesIO(source), data_only=False)
        updated = load_workbook(BytesIO(revised), data_only=False)
        if original.sheetnames != updated.sheetnames:
            return False
        allowed = {(e.sheet_name, e.cell): e for e in edits}
        for old_sheet, new_sheet in zip(original.worksheets, updated.worksheets):
            if (str(old_sheet.merged_cells) != str(new_sheet.merged_cells)
                    or str(old_sheet.print_area) != str(new_sheet.print_area)
                    or old_sheet.max_row != new_sheet.max_row
                    or old_sheet.max_column != new_sheet.max_column):
                return False
            for row in old_sheet:
                for old_cell in row:
                    new_cell = new_sheet[old_cell.coordinate]
                    edit = allowed.get((old_sheet.title, old_cell.coordinate))
                    if edit:
                        if old_cell.value != edit.previous_value or new_cell.value != edit.replacement_value:
                            return False
                    elif old_cell.value != new_cell.value:
                        return False
                    if (old_cell.style_id != new_cell.style_id
                            or old_cell.number_format != new_cell.number_format
                            or (old_cell.data_type == "f") != (new_cell.data_type == "f")):
                        return False
        return True


class FakePdfRenderer:
    def __init__(self, value=PDF):
        self.value = value
        self.calls = 0

    def render_xlsx_pdf(self, data):
        self.calls += 1
        return self.value


def execute(source=None, editor=None, verifier=None, renderer=None, **updates):
    args = dict(
        source=source if source is not None else synthetic_source(),
        edits=EDITS,
        editor=editor if editor is not None else TrustedEditor(),
        verifier=verifier if verifier is not None else IndependentVerifier(),
        renderer=renderer if renderer is not None else FakePdfRenderer(),
        workspace_ref=WORKSPACE,
        run_ref=RUN,
        revised_artifact_id="quote_revised_3933",
        pdf_artifact_id="quote_pdf_3933",
        lineage_id="quote_lineage_3933",
        now=NOW,
    )
    args.update(updates)
    return revise_and_render_local_xlsx_pdf(**args)


class ReviewedXlsxRevisionContractTests(unittest.TestCase):
    def test_real_xlsx_edit_semantics_and_canonical_pdf_lineage(self):
        source = synthetic_source()
        original_bytes, original_digest = source.content, sha256(source.content).hexdigest()
        result = execute(source=source)
        assert isinstance(result, LocalRevisedXlsxPdf)
        assert source.content == original_bytes
        assert sha256(source.content).hexdigest() == original_digest
        actual = load_workbook(BytesIO(result.revised_bytes), data_only=False).active
        assert actual["A1"].value == "New School"
        assert actual["B2"].value == 99_200_000
        assert actual["B3"].value == "=B2*0.1"
        assert actual["B4"].value == "=B2+B3"
        assert actual["A1"].font.bold is True
        assert "A5:C5" in str(actual.merged_cells)
        assert result.revised.integrity_ref == sha256(result.revised_bytes).hexdigest()
        assert result.pdf.integrity_ref == sha256(PDF).hexdigest()
        assert result.lineage.source.artifact_id == source.record.artifact_id
        assert result.lineage.source.integrity_ref == original_digest
        assert result.lineage.working_artifact_id == result.revised.artifact_id
        assert result.lineage.output_artifact_ids == (result.pdf.artifact_id,)
        for record, content in (
            (result.revised, result.revised_bytes), (result.pdf, result.pdf_bytes)
        ):
            ref = LineageArtifactRef(record.artifact_id, record.integrity_ref)
            material = result.resolve_artifact_material(ref, workspace_ref=WORKSPACE, run_ref=RUN)
            assert material is not None
            assert material.content == content
            assert result.resolve_artifact_material(
                ref, workspace_ref="foreign_workspace", run_ref=RUN
            ) is None
            assert result.resolve_artifact_material(
                ref, workspace_ref=WORKSPACE, run_ref="foreign_run"
            ) is None
        projection = result.public_projection()
        assert projection["drive_write_authorized"] is False
        assert projection["pdf_fidelity_attested"] is False
        assert projection["raw_document_bytes"] is False
        assert not PRODUCTION_REVISION_HOST_COMPOSED
        assert not PRODUCTION_DRIVE_WRITE_AUTHORIZED
        assert "Existing School" not in str(projection)


    def test_unknown_source_or_wrong_run_denied_before_edit(self):
        editor = TrustedEditor()
        source = synthetic_source()
        for kwargs in (
            {"workspace_ref": "foreign_workspace"},
            {"run_ref": "other_run"},
            {"source": replace(source, content=source.content + b"tampered")},
            {"source": replace(source, route=replace(source.route, reason="model_claim"))},
        ):
            with self.assertRaises(ContractError):
                execute(editor=editor, **kwargs)
        assert editor.calls == 0


    def test_verification_refusal_prevents_pdf_and_registration(self):
        editor = TrustedEditor()
        verifier = IndependentVerifier(accepted=False)
        renderer = FakePdfRenderer()
        with self.assertRaisesRegex(ContractError, "requested-only"):
            execute(editor=editor, verifier=verifier, renderer=renderer)
        assert editor.calls == verifier.calls == 1
        assert renderer.calls == 0


    def test_duplicate_or_unreviewed_edits_rejected_before_execution(self):
        editor = TrustedEditor()
        for edits in (
            (),
            EDITS + (EDITS[0],),
            EDITS * (MAX_REVISIONS // 2 + 1),
            ("not an edit",),
        ):
            with self.assertRaises(ContractError):
                execute(edits=edits, editor=editor)
        assert editor.calls == 0
        for args in (
            ("Quote", "B3", "=B2*0.1", "100"),
            ("Quote", "A1", "School", "=EVIL()"),
            ("Quote", "A1", "School", "-EVIL()"),
            ("Quote", "A1", "School", "School"),
            ("Quote", "A0", 1, 2),
            ("Quote", "A1", True, 2),
        ):
            with self.assertRaises(ContractError):
                ReviewedCellRevision(*args)


    def test_invalid_editor_or_invalid_pdf_fail_closed(self):
        class Unchanged:
            def revise_xlsx_copy(self, data, edits):
                return data

        class Corrupt:
            def revise_xlsx_copy(self, data, edits):
                return b"not an xlsx"

        for editor in (Unchanged(), Corrupt(), object()):
            with self.assertRaises(ContractError):
                execute(editor=editor)
        for pdf in (b"", b"no pdf", b"%PDF-1.4\nmissing eof"):
            with self.assertRaises(ContractError):
                execute(renderer=FakePdfRenderer(pdf))


    def test_duplicate_artifact_id_rejected_pre_execution(self):
        editor = TrustedEditor()
        with self.assertRaises(ContractError):
            execute(editor=editor, revised_artifact_id="quote_source_3933")
        assert editor.calls == 0


    def test_tampered_output_material_refused(self):
        result = execute()
        ref = LineageArtifactRef(result.pdf.artifact_id, result.pdf.integrity_ref)
        assert replace(result, pdf_bytes=PDF + b"spoof").resolve_artifact_material(
            ref, workspace_ref=WORKSPACE, run_ref=RUN
        ) is None
        assert replace(result, pdf=replace(result.pdf, workspace_ref="foreign_workspace")).resolve_artifact_material(
            ref, workspace_ref=WORKSPACE, run_ref=RUN
        ) is None
        assert replace(result, lineage=replace(result.lineage, run_ref="other_run")).resolve_artifact_material(
            ref, workspace_ref=WORKSPACE, run_ref=RUN
        ) is None


    def test_stale_formula_cache_never_called_verified_pdf(self):
        result = execute()
        assert result.public_projection()["pdf_fidelity_attested"] is False
        # Formula AST preserved; no claim of Excel recalc/cache or printed values.
        wb = load_workbook(BytesIO(result.revised_bytes), data_only=False)
        assert wb.active["B3"].value == "=B2*0.1"
