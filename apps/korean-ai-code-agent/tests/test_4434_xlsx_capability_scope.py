"""#4434: the boundary between the ordinary Python XLSX path and the Office path.

CENTRAL holds ``classify_xlsx(local_only=True)`` as the authority: REJECT is
refused, ``office_ephemeral_or_desktop_fallback`` stays on the existing Office
route, and only a Python-eligible workbook may reach the copy editor. So these
cases use two kinds of real bytes — a minimal package that carries no styles
part, no theme, no merge and no formula, and openpyxl-authored books, which are
fidelity-sensitive by construction — and each one states the route outcome the PC
lane reports to the web task.
"""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import unittest
import zipfile

from openpyxl import Workbook, load_workbook

from padiem_ai_core.document_normalization import (
    extract_binary_document, read_xlsx_table,
)

from kagent.local_xlsx_python_copy_editor import (
    IndependentZipXlsxRevisionVerifier, OfficeFidelityRequired, PythonXlsxValueCopyEditor,
)
from kagent.contracts import ContractError
from kagent.local_xlsx_revision_pdf import ReviewedCellRevision
from kagent.xlsx_fidelity_route import XlsxRoute, classify_xlsx

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
OFFICE_ONLY_PARTS = (
    "xl/vbaProject.bin",
    "xl/externalLinks/externalLink1.xml",
    "xl/pivotCache/pivotCacheDefinition1.xml",
    "xl/charts/chart1.xml",
    "xl/comments1.xml",
    "customxml/item1.xml",
)

EDITOR = PythonXlsxValueCopyEditor()
VERIFIER = IndependentZipXlsxRevisionVerifier()

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def build(*, sheets: int = 1, formula: bool = False, styles: bool = True) -> bytes:
    """An Excel-shaped workbook: openpyxl always makes this fidelity-sensitive."""
    book = Workbook()
    first = book.active
    first.title = "Quote"
    first["A1"] = "기존 항목"
    first["B2"] = 88_000_000
    first["C2"] = "text kept"
    if styles:
        first["A5"] = "merged note"
        first.merge_cells("A5:C5")
    if formula:
        first["B3"] = "=B2*0.1"
    for index in range(2, sheets + 1):
        sheet = book.create_sheet(f"Sheet{index}")
        sheet["A1"] = f"시트 {index} 값"
        sheet["B2"] = index * 1000
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def eligible_bytes(*, sheets: int = 1, rows: int = 0) -> bytes:
    """A workbook the classifier admits to the Python path.

    No styles part, no theme, no calcChain, no ``s`` attribute, no merge element:
    exactly the structure that is not fidelity-sensitive.
    """
    names = ["Quote"] + [f"Sheet{index}" for index in range(2, sheets + 1)]
    overrides, rels, sheet_entries = [], [], []
    parts = {}
    for index, title in enumerate(names, start=1):
        part = f"xl/worksheets/sheet{index}.xml"
        overrides.append(
            f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        )
        rels.append(
            f'<Relationship Id="rId{index + 3}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{index}.xml"/>'
        )
        sheet_entries.append(f'<sheet name="{title}" sheetId="{index}" r:id="rId{index + 3}"/>')
        if index == 1 and rows:
            body = "".join(
                f'<row r="{row}"><c r="A{row}" t="inlineStr"><is><t>행 {row}</t></is></c>'
                f'<c r="B{row}"><v>{row}</v></c></row>'
                for row in range(1, rows + 1)
            )
            extent = f"A1:B{rows}"
        else:
            body = ('<row r="1"><c r="A1" t="inlineStr"><is><t>기존 항목</t></is></c>'
                    '<c r="B1" t="inlineStr"><is><t>금액</t></is></c></row>'
                    '<row r="2"><c r="B2"><v>88000000</v></c>'
                    '<c r="C2" t="inlineStr"><is><t>text kept</t></is></c></row>')
            extent = "A1:C2"
            if index > 1:
                body = (f'<row r="1"><c r="A1" t="inlineStr"><is><t>시트 {index} 값</t></is></c>'
                        f'<c r="B1" t="inlineStr"><is><t>수치</t></is></c></row>'
                        f'<row r="2"><c r="B2"><v>{index * 1000}</v></c></row>')
                extent = "A1:B2"
        parts[part] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                       f'<worksheet xmlns="{MAIN}">'
                       f'<dimension ref="{extent}"/><sheetData>{body}</sheetData></worksheet>')
    parts["[Content_Types].xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        + "".join(overrides) + "</Types>"
    )
    parts["_rels/.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PKG_REL}">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    parts["xl/workbook.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<workbook xmlns="{MAIN}" xmlns:r="{REL}"><sheets>'
        + "".join(sheet_entries) + "</sheets></workbook>"
    )
    parts["xl/_rels/workbook.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PKG_REL}">' + "".join(rels) + "</Relationships>"
    )
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, text in parts.items():
            archive.writestr(name, text)
    return buffer.getvalue()


def with_part(data: bytes, name: str, payload: bytes = b"\x00\x01binary") -> bytes:
    source = zipfile.ZipFile(BytesIO(data), mode="r")
    output = BytesIO()
    with zipfile.ZipFile(output, mode="w") as archive:
        for info in source.infolist():
            archive.writestr(info, source.read(info.filename))
        archive.writestr(name, payload)
    source.close()
    return output.getvalue()


def edit(sheet: str, cell: str, previous, replacement) -> ReviewedCellRevision:
    return ReviewedCellRevision(
        sheet_name=sheet, cell=cell, previous_value=previous,
        replacement_value=replacement,
    )


def changed_parts(source: bytes, revised: bytes) -> list[str]:
    before = zipfile.ZipFile(BytesIO(source))
    after = zipfile.ZipFile(BytesIO(revised))
    return [name for name in before.namelist() if before.read(name) != after.read(name)]


class XlsxRouteAuthorityTests(unittest.TestCase):
    """The classifier decides the route; the editor cannot claim a file back."""

    def test_eligibility_is_reported_separately_from_edit_capability(self) -> None:
        """A styled workbook is refused even though the editor could handle it.

        This is the mismatch CENTRAL asked to see pinned: openpyxl's output is
        structurally simple enough for a literal value copy, and it is still the
        Office route's, because the styles part and the style indices make it
        fidelity-sensitive. Admitting it here would be a silent downgrade of the
        route on capability grounds.
        """
        source = build()

        self.assertIs(classify_xlsx(source, local_only=True).route,
                      XlsxRoute.OFFICE_FALLBACK)
        with self.assertRaises(OfficeFidelityRequired) as refused:
            EDITOR.revise_xlsx_copy(source, (edit("Quote", "B2", 88_000_000, 1),))
        self.assertIn("fidelity_sensitive_workbook_features", str(refused.exception))

    def test_the_minimal_package_is_the_python_eligible_case(self) -> None:
        source = eligible_bytes()
        decision = classify_xlsx(source, local_only=True)

        self.assertIs(decision.route, XlsxRoute.NATIVE_CANDIDATE_REQUIRES_PROOF)
        self.assertEqual(decision.execution_target, "padiem_desktop_local_runner")
        self.assertIs(decision.may_automatically_convert, False)

    def test_merged_cells_send_the_workbook_to_the_office_path(self) -> None:
        merged = eligible_bytes()
        merged = _add_merge(merged)

        self.assertIs(classify_xlsx(merged, local_only=True).route,
                      XlsxRoute.OFFICE_FALLBACK)
        with self.assertRaises(OfficeFidelityRequired):
            EDITOR.revise_xlsx_copy(merged, (edit("Quote", "B2", 88_000_000, 1),))


class XlsxCapabilityScopeTests(unittest.TestCase):
    def test_value_edit_in_a_later_sheet_leaves_the_other_sheets_alone(self) -> None:
        source = eligible_bytes(sheets=3)
        edits = (edit("Sheet2", "B2", 2000, 777_000),)
        revised = EDITOR.revise_xlsx_copy(source, edits)

        self.assertIs(VERIFIER.verify_only_reviewed_edits(source, revised, edits), True)
        self.assertEqual(changed_parts(source, revised), ["xl/worksheets/sheet2.xml"])
        book = load_workbook(BytesIO(revised))
        self.assertEqual(book["Sheet2"]["B2"].value, 777_000)
        self.assertEqual(book["Sheet3"]["B2"].value, 3000)
        self.assertEqual(book["Quote"]["B2"].value, 88_000_000)
        self.assertEqual(book["Sheet2"]["A1"].value, "시트 2 값")
        book.close()

    def test_korean_text_and_multi_cell_revisions_round_trip(self) -> None:
        source = eligible_bytes(sheets=3)
        edits = (
            edit("Quote", "C2", "text kept", "  최종 견적  "),
            edit("Quote", "B2", 88_000_000, -3_500),
            edit("Sheet2", "A1", "시트 2 값", "rename confirmed"),
        )
        revised = EDITOR.revise_xlsx_copy(source, edits)
        self.assertIs(VERIFIER.verify_only_reviewed_edits(source, revised, edits), True)

        book = load_workbook(BytesIO(revised))
        self.assertEqual(book["Quote"]["C2"].value, "  최종 견적  ")
        self.assertEqual(book["Quote"]["B2"].value, -3_500)
        self.assertEqual(book["Sheet2"]["A1"].value, "rename confirmed")
        book.close()

        flat = extract_binary_document(name="book.xlsx", media_type=XLSX_MIME, payload=revised)
        self.assertIn("[Quote!B2] -3500", flat.text)
        # Leading and trailing spaces survive because the replacement is
        # written with xml:space="preserve".
        self.assertIn("[Quote!C2]   최종 견적  ", flat.text)
        self.assertEqual(read_xlsx_table(
            name="book.xlsx", media_type=XLSX_MIME, payload=revised,
        ).sheet_title, "Quote")

    def test_many_rows_are_edited_without_touching_the_rest(self) -> None:
        source = eligible_bytes(rows=200)
        self.assertIs(classify_xlsx(source, local_only=True).route,
                      XlsxRoute.NATIVE_CANDIDATE_REQUIRES_PROOF)

        edits = (edit("Quote", "B150", 150, 1_500_000),)
        revised = EDITOR.revise_xlsx_copy(source, edits)
        self.assertIs(VERIFIER.verify_only_reviewed_edits(source, revised, edits), True)
        checked = load_workbook(BytesIO(revised))["Quote"]
        self.assertEqual(checked["B150"].value, 1_500_000)
        self.assertEqual(checked["A149"].value, "행 149")
        self.assertEqual(checked["B200"].value, 200)
        checked.parent.close()

    def test_workbooks_the_office_path_must_handle_are_routed_away(self) -> None:
        for extra_part in OFFICE_ONLY_PARTS:
            with self.subTest(extra_part=extra_part):
                source = with_part(eligible_bytes(), extra_part)
                self.assertIs(
                    classify_xlsx(source, local_only=True).route, XlsxRoute.OFFICE_FALLBACK,
                )
                with self.assertRaises(OfficeFidelityRequired):
                    EDITOR.revise_xlsx_copy(source, (edit("Quote", "B2", 88_000_000, 1),))

    def test_formula_and_calculation_chain_never_get_a_stale_result(self) -> None:
        formula = with_part(eligible_bytes(), "xl/calcChain.xml", b"<calcChain/>")
        with self.assertRaises(OfficeFidelityRequired):
            EDITOR.revise_xlsx_copy(formula, (edit("Quote", "B2", 88_000_000, 1),))

        inline_formula = eligible_bytes()
        inline_formula = _add_formula_cell(inline_formula)
        self.assertIs(classify_xlsx(inline_formula, local_only=True).route,
                      XlsxRoute.OFFICE_FALLBACK)
        with self.assertRaises(OfficeFidelityRequired):
            EDITOR.revise_xlsx_copy(inline_formula, (edit("Quote", "B3", 8_800_000, 1),))

        styled_formula = build(formula=True)
        with self.assertRaises(OfficeFidelityRequired):
            EDITOR.revise_xlsx_copy(styled_formula, (edit("Quote", "B2", 88_000_000, 1),))

    def test_untrusted_or_unsupported_input_is_refused_before_routing(self) -> None:
        self.assertIs(classify_xlsx(b"PK\x03\x04", local_only=True).route, XlsxRoute.REJECT)
        with self.assertRaises(ContractError):
            EDITOR.revise_xlsx_copy(b"PK\x03\x04garbage", (edit("Quote", "B2", 1, 2),))
        with self.assertRaises(ContractError):
            EDITOR.revise_xlsx_copy(eligible_bytes(), (edit("Quote", "K7", 0, 1),))
        with self.assertRaises(ContractError):
            EDITOR.revise_xlsx_copy(eligible_bytes(), (edit("Absent", "B2", 88_000_000, 1),))

    def test_literal_only_revision_contract_refuses_float_and_code(self) -> None:
        with self.assertRaises(ContractError):
            edit("Quote", "D3", 12.5, 14.75)
        with self.assertRaises(ContractError):
            edit("Quote", "C2", "text kept", "=HYPERLINK(\"https://x\")")
        with self.assertRaises(ContractError):
            edit("Quote", "B2", 88_000_000, True)

    def test_original_bytes_are_untouched_by_a_revision(self) -> None:
        source = eligible_bytes()
        digest = sha256(source).hexdigest()
        EDITOR.revise_xlsx_copy(source, (edit("Quote", "B2", 88_000_000, 42),))
        self.assertEqual(sha256(source).hexdigest(), digest)


def _rewrite_part(data: bytes, name: str, transform) -> bytes:
    source = zipfile.ZipFile(BytesIO(data), mode="r")
    output = BytesIO()
    with zipfile.ZipFile(output, mode="w") as archive:
        for info in source.infolist():
            text = source.read(info.filename)
            if info.filename == name:
                text = transform(text.decode("utf-8"))
            archive.writestr(info, text)
    source.close()
    return output.getvalue()


def _add_merge(data: bytes) -> bytes:
    return _rewrite_part(
        data, "xl/worksheets/sheet1.xml",
        lambda text: text.replace("</sheetData>",
                                  '</sheetData><mergeCells count="1">'
                                  '<mergeCell ref="A5:C5"/></mergeCells>'),
    )


def _add_formula_cell(data: bytes) -> bytes:
    return _rewrite_part(
        data, "xl/worksheets/sheet1.xml",
        lambda text: text.replace("</row></sheetData>",
                                  '<c r="B3"><f>B2*0.1</f><v>8800000</v></c>'
                                  "</row></sheetData>"),
    )


if __name__ == "__main__":
    unittest.main()
