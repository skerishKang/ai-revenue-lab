"""The exported files are safe to open and carry the values they were given."""

from __future__ import annotations

import csv
from pathlib import Path

from app.exporter import neutralize, write_rows, write_workbook


def _read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [list(row) for row in csv.reader(handle)]


def test_a_formula_lookalike_is_neutralised_before_it_is_written(tmp_path) -> None:
    target = write_rows(
        tmp_path / "items.csv",
        ["상품명"],
        [[neutralize("=1+1")], ["정상품"]],
    )

    rows = _read_csv(target)
    assert rows[1] == ["'=1+1"]
    assert rows[2] == ["정상품"]


def test_neutralize_leaves_ordinary_values_alone() -> None:
    assert neutralize("6205-2RS") == "6205-2RS"
    assert neutralize("볼 베어링") == "볼 베어링"
    assert neutralize("") == ""


def test_an_empty_value_is_written_empty_and_never_defaulted(tmp_path) -> None:
    target = write_rows(tmp_path / "items.csv", ["품목코드", "매입가"], [["A-1", ""]])

    rows = _read_csv(target)
    assert rows[1] == ["A-1", ""]


def test_the_csv_is_written_with_a_bom_so_excel_reads_korean_headers(tmp_path) -> None:
    target = write_rows(tmp_path / "items.csv", ["상품명"], [["볼 베어링"]])

    assert target.read_bytes().startswith(b"\xef\xbb\xbf")


def test_the_workbook_and_the_csv_carry_the_same_values(tmp_path) -> None:
    from openpyxl import load_workbook

    headers = ["품목코드", "상품명", "매입가"]
    rows = [["BRG-6205-01", "볼 베어링 6205-2RS", "8500"], ["", "알루미늄 40x40", ""]]

    csv_path = write_rows(tmp_path / "items.csv", headers, rows)
    workbook_path = write_workbook(
        tmp_path / "items.xlsx",
        sheets=[("품목마스터", headers, rows)],
        summary={"전체행수": 2},
    )

    workbook = load_workbook(workbook_path)
    sheet = workbook["품목마스터"]
    workbook_rows = [
        ["" if cell.value is None else str(cell.value) for cell in line]
        for line in sheet.iter_rows()
    ]
    workbook.close()

    assert workbook_rows == _read_csv(csv_path)
