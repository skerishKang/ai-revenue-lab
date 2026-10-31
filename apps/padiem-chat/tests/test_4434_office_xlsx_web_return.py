"""#4434: the approved PC workbook comes back to the web user it belongs to.

The Office chunk lane already stages returned bytes in the private Broker store
and the server already rebuilds and verifies them. What was missing is the
owner-facing web seam: a query that reports which outputs one of the owner's own
completed runs has, and the download of the revised workbook itself.

These tests pin those two routes: the owner gets exactly the staged bytes under
the artifact's own verified name, a caller who is not the owner learns nothing,
an unconfigured staging lane says "unconfigured" instead of pretending the run is
empty, and a stream that fails integrity never reaches a browser. The handlers
are called the way the router calls them, on an app state that carries the same
settings, history store and staging source the worker sets.
"""
from __future__ import annotations

import asyncio
from base64 import b64encode
from hashlib import sha256
import json
from pathlib import Path
import re
from urllib.parse import quote

import pytest
from starlette.datastructures import MutableHeaders, State
from starlette.requests import Request

from app.auth import SESSION_COOKIE, create_session_token
from app.claw_local_office_binary_reader import OFFICE_CHUNK_BYTES, VerifiedOfficeBinary
from app.claw_office_preview_routes import (
    OFFICE_OUTPUTS_QUERY_PATH, OFFICE_PDF_PREVIEW_PATH, OFFICE_XLSX_DOWNLOAD_PATH,
    claw_office_outputs_query, claw_office_pdf_preview, claw_office_xlsx_download,
)
from app.config import Settings

OWNER = "usr_4434_office_owner"
FOREIGN = "usr_4434_office_other"
RUN = "run_4434_office_return"

XLSX_NAME = "수정된 견적서.xlsx"
PDF_NAME = "verified.pdf"
PDF_MEDIA = "application/pdf"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ORIGIN = "https://chat.example.test"

# A zip container header plus enough body to have to arrive in more than one
# chunk; the real reassembly through the composition is the last test's job.
XLSX = b"PK\x03\x04" + bytes(range(256)) * 300
XLSX_PARTS = (len(XLSX) + OFFICE_CHUNK_BYTES - 1) // OFFICE_CHUNK_BYTES


def real_workbook(rows: int) -> bytes:
    """A genuine OOXML package, big enough that the lane has to chunk it.

    The reader runs the same untrusted-input classifier on the rebuilt bytes that
    the PC route gate uses, so the composition test cannot use a zip-shaped guess.
    """
    import io

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "견적"
    sheet["A1"], sheet["B1"] = "항목", "금액"
    for index in range(2, rows + 2):
        sheet.cell(row=index, column=1, value=f"항목 {index}")
        sheet.cell(row=index, column=2, value=index * 1000)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()

PDF = b"%PDF-1.4\n" + b"X" * (OFFICE_CHUNK_BYTES + 40) + b"\n%%EOF\n"

HANDLERS = {
    OFFICE_XLSX_DOWNLOAD_PATH: claw_office_xlsx_download,
    OFFICE_OUTPUTS_QUERY_PATH: claw_office_outputs_query,
    OFFICE_PDF_PREVIEW_PATH: claw_office_pdf_preview,
}


class StagedSource:
    """The staging lane as the composition would hand it to the routes."""

    configured = True

    def __init__(self):
        self.contents = {"xlsx": XLSX, "pdf": PDF}
        self.names = {"xlsx": XLSX_NAME, "pdf": PDF_NAME}
        self.calls = []

    async def read_staged_office_output(self, *, owner_id, run_ref, kind):
        self.calls.append((owner_id, run_ref, kind))
        if owner_id != OWNER or run_ref != RUN:
            raise ValueError("PRIVATE_DEVICE_OR_OWNER_NOT_AVAILABLE")
        data = self.contents.get(kind)
        if data is None:
            raise ValueError("office_artifact_unavailable")
        return VerifiedOfficeBinary(
            artifact_id=f"artifact_office_4434_{kind}",
            filename=self.names[kind], kind=kind,
            media_type=PDF_MEDIA if kind == "pdf" else XLSX_MEDIA,
            integrity_ref=sha256(data).hexdigest(), content=data,
        )


class UnconfiguredSource(StagedSource):
    configured = False


class SilentSource:
    """The served Chat Worker default: no staging lane composed at all."""


class Store:
    async def get_user(self, user_id):
        return None

    async def list_projects(self, user_id):
        return []


def settings() -> Settings:
    return Settings(
        session_secret="office-4434-web-return-tests",
        auth_mode="mock", public_base_url=ORIGIN,
    )


def request(path, source, *, run=RUN, owner=OWNER, state_settings=None):
    """The Request the router would hand the handler for this call."""
    target = path.replace("{run_id}", run)
    state = State()
    state.settings = state_settings or settings()
    state.history_store = Store()
    state.local_task_result_source = source
    headers = MutableHeaders()
    if owner is not None:
        headers["Cookie"] = SESSION_COOKIE + "=" + create_session_token(
            state.settings, owner)
    return Request({
        "type": "http", "http_version": "1.1", "scheme": "https",
        "method": "GET", "path": target, "raw_path": target.encode(),
        "query_string": b"", "root_path": "",
        "headers": [(k.lower().encode(), v.encode()) for k, v in headers.items()],
        "client": ("127.0.0.1", 50000), "server": (b"chat.example.test", 443),
        "app": _AppState(state), "path_params": {"run_id": run},
    })


def call(path, source, **kwargs):
    return asyncio.run(HANDLERS[path](request(path, source, **kwargs)))



class _AppState:
    def __init__(self, state):
        self.state = state


def body(response):
    return json.loads(response.body.decode("utf-8"))


def test_owner_downloads_the_exact_staged_workbook_under_its_own_name():
    source = StagedSource()
    response = call(OFFICE_XLSX_DOWNLOAD_PATH, source)
    assert response.status_code == 200
    assert response.body == XLSX
    assert sha256(response.body).hexdigest() == sha256(XLSX).hexdigest()
    assert XLSX_PARTS > 1, "the fixture has to be a chunked artifact"
    assert response.headers["content-type"] == XLSX_MEDIA
    assert response.headers["content-disposition"] == (
        f"attachment; filename*=UTF-8''{quote(XLSX_NAME, safe='')}"
    )
    assert response.headers["cache-control"].startswith("private, no-store")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["vary"] == "Cookie"
    assert response.headers["content-length"] == str(len(XLSX))
    # The kind is fixed by the route: a browser cannot ask for another one.
    assert source.calls == [(OWNER, RUN, "xlsx")]


def test_query_lists_every_available_output_and_carries_no_bytes():
    source = StagedSource()
    response = call(OFFICE_OUTPUTS_QUERY_PATH, source)
    assert response.status_code == 200
    payload = body(response)
    assert payload["ok"] is True and payload["run_id"] == RUN
    assert payload["private_staging_only"] is True
    assert payload["raw_bytes_in_projection"] is False
    assert {item["media_type"] for item in payload["outputs"]} == {PDF_MEDIA, XLSX_MEDIA}
    for item in payload["outputs"]:
        assert set(item) == {
            "artifact_id", "filename", "media_type", "size_bytes", "integrity_ref",
            "private_staging_only", "drive_write_approval", "raw_bytes_in_projection",
        }
        assert item["drive_write_approval"] is False
        assert len(item["integrity_ref"]) == 64
    assert b"PK" not in response.body and b"data_b64" not in response.body
    assert {call_[2] for call_ in source.calls} == {"pdf", "xlsx"}


def test_only_the_staged_kinds_are_listed():
    source = StagedSource()
    source.contents.pop("pdf")
    response = call(OFFICE_OUTPUTS_QUERY_PATH, source)
    assert response.status_code == 200
    assert [item["media_type"] for item in body(response)["outputs"]] == [XLSX_MEDIA]


def test_no_session_never_reaches_the_private_staging():
    source = StagedSource()
    for path in (OFFICE_XLSX_DOWNLOAD_PATH, OFFICE_OUTPUTS_QUERY_PATH):
        response = call(path, source, owner=None)
        assert response.status_code == 401, path
        assert body(response)["error"]["code"] == "unauthorized"
    assert source.calls == []


def test_another_owner_and_an_unknown_run_give_the_identical_answer():
    source = StagedSource()
    foreign = call(OFFICE_XLSX_DOWNLOAD_PATH, source, owner=FOREIGN)
    unknown = call(OFFICE_XLSX_DOWNLOAD_PATH, source, run="run_nobody_owns")
    assert foreign.status_code == unknown.status_code == 404
    assert body(foreign) == body(unknown)
    assert b"PRIVATE_DEVICE" not in foreign.body
    # The query keeps the same discipline: an empty list would confirm that the
    # run exists and that the caller owns it.
    query = call(OFFICE_OUTPUTS_QUERY_PATH, source, owner=FOREIGN)
    assert query.status_code == 404
    # A refusal carries only the bounded code: no outputs list, so no caller
    # learns that the run exists and just happens to have nothing staged.
    assert set(body(query)) == {"ok", "error"}


def test_an_unconfigured_staging_lane_says_unconfigured_not_empty():
    cases = (
        (OFFICE_XLSX_DOWNLOAD_PATH, "office_xlsx_source_unconfigured"),
        (OFFICE_OUTPUTS_QUERY_PATH, "office_outputs_source_unconfigured"),
        (OFFICE_PDF_PREVIEW_PATH, "office_pdf_source_unconfigured"),
    )
    for path, code in cases:
        for source in (UnconfiguredSource(), SilentSource()):
            response = call(path, source)
            assert response.status_code == 503, (path, type(source).__name__)
            assert body(response)["error"]["code"] == code


def test_a_body_that_is_not_the_declared_container_is_not_forwarded():
    # The reader checks the whole-file digest, so a self-consistent stream with
    # the wrong container has to be stopped by the route's own signature test.
    source = StagedSource()
    source.contents["xlsx"] = b"not a zip at all"
    response = call(OFFICE_XLSX_DOWNLOAD_PATH, source)
    assert response.status_code == 503
    assert body(response)["error"]["code"] == "office_xlsx_invalid"
    assert b"PK" not in response.body


def test_a_truncated_container_below_the_floor_is_not_offered_as_a_download():
    # The reader owns the name and extension contract; the route owns the last
    # check on what actually goes to the browser: a declared-size stream that
    # came back a stub is refused, not served.
    source = StagedSource()
    source.contents["xlsx"] = b"PK" + b"Z" * 10
    assert call(OFFICE_XLSX_DOWNLOAD_PATH, source).status_code == 503


def test_a_run_id_that_is_not_a_run_id_is_refused_before_the_source():
    source = StagedSource()
    for bad in ("..", "run/../x", "a" * 200):
        assert call(OFFICE_XLSX_DOWNLOAD_PATH, source, run=bad).status_code == 400, bad
    assert source.calls == []


def test_auth_mode_off_does_not_serve_any_owner_output():
    source = StagedSource()
    off = Settings(session_secret="office-4434", auth_mode="off", public_base_url=ORIGIN)
    for path in (OFFICE_XLSX_DOWNLOAD_PATH, OFFICE_OUTPUTS_QUERY_PATH):
        assert call(path, source, state_settings=off).status_code == 401
    assert source.calls == []


def test_the_web_return_added_no_authority_and_enabled_nothing():
    from app.claw_local_office_binary_reader import PRODUCTION_OFFICE_PRIVATE_READ_WIRED
    from app.claw_office_preview_routes import (
        OFFICE_PDF_PREVIEW_DRIVE_WRITE, OFFICE_PDF_PREVIEW_NEW_APPROVAL_AUTHORITY,
        OFFICE_XLSX_DOWNLOAD_DRIVE_WRITE, OFFICE_XLSX_DOWNLOAD_NEW_APPROVAL_AUTHORITY,
    )

    assert PRODUCTION_OFFICE_PRIVATE_READ_WIRED is False
    assert OFFICE_PDF_PREVIEW_NEW_APPROVAL_AUTHORITY is False
    assert OFFICE_PDF_PREVIEW_DRIVE_WRITE is False
    assert OFFICE_XLSX_DOWNLOAD_NEW_APPROVAL_AUTHORITY is False
    assert OFFICE_XLSX_DOWNLOAD_DRIVE_WRITE is False
    # The committed config still leaves the chunk gate absent, so the staging
    # store composes to None and these routes answer 503 in Production.
    root = Path(__file__).resolve().parents[3]
    chat = (root / "apps" / "padiem-chat" / "wrangler.toml").read_text(encoding="utf-8")
    broker = (root / "packages" / "padiem-control-plane" /
              "wrangler.local-agent-broker.jsonc").read_text(encoding="utf-8")
    assert "LOCAL_AGENT_OFFICE_CHUNK_TRANSFER_ENABLED" not in chat
    assert "LOCAL_AGENT_OFFICE_CHUNK_TRANSFER_ENABLED" not in broker


def test_both_routes_are_registered_and_the_owner_read_surface_stays_bounded():
    factory = (Path(__file__).resolve().parents[1] / "app" / "app_factory.py").read_text(
        encoding="utf-8"
    )
    for template, handler in (
        ("OFFICE_XLSX_DOWNLOAD_PATH", "claw_office_xlsx_download"),
        ("OFFICE_OUTPUTS_QUERY_PATH", "claw_office_outputs_query"),
    ):
        assert len(re.findall(
            rf'Route\({template}, {handler}, methods=\["GET"\]\)', factory)) == 1
    # Exactly three owner-facing Office output GET routes: pdf, xlsx, outputs.
    assert len(re.findall(r"Route\(OFFICE_[A-Z_]+PATH,", factory)) == 3
    # The web return added no device-facing surface: the broker device route set
    # the B54 gate workflow pins still has exactly one Office entry and it is the
    # staging write the Connector already uses, never a byte download.
    from local_agent_broker_device_http import _DEVICE_HTTP_ROUTES

    routes = set(_DEVICE_HTTP_ROUTES)
    assert "/office-part" in routes
    assert [route for route in routes if "office" in route] == ["/office-part"]
    assert not [route for route in routes if route.startswith("/api/")]


class RunHistory:
    """The three owner-scoped reads the composition makes, in memory.

    The D1 shapes behind them are pinned by the composition suite; what is proved
    here is the seam between the two web routes and that composition, so the run
    row, the stored correlation and the terminal fact are supplied the way the
    server supplies them and every byte still goes through the real reader.
    """

    OWNER = "usr_4434_composition_owner"
    OTHER = "usr_4434_composition_other"
    RUN = "run_4434_composition"
    WORKSPACE = "ws_4434_composition"
    CONVERSATION = "chat_" + "4434" * 8
    CORRELATION = {
        "command_id": "cmd_4434_office_return",
        "tool_request_ref": "tool_4434_office_return",
        "request_id": "request_4434_office_return",
        "revision_ref": "revision_4434_office_return",
        "evidence_ref": "evidence_4434_office_return",
        "request_fingerprint": "b" * 64,
    }

    def __init__(self, *, completed=True):
        self.completed = completed

    async def get_claw_run(self, user_id, run_id):
        if user_id != self.OWNER or run_id != self.RUN:
            raise LookupError("not found")
        return {
            "run_id": self.RUN, "conversation_id": self.CONVERSATION,
            "workspace_id": self.WORKSPACE,
            "status": "completed" if self.completed else "running",
        }

    async def verify_owner_conversation_artifacts(self, user_id, conversation_id):
        return user_id == self.OWNER and conversation_id == self.CONVERSATION

    async def get_local_task_correlation(self, user_id, run_id):
        if user_id != self.OWNER or run_id != self.RUN:
            raise LookupError("not bound")
        return dict(self.CORRELATION)


class TerminalPort:
    """The canonical Broker terminal fact and the private part stream."""

    def __init__(self, content, *, drop_last=False):
        self.content = content
        self.parts = -(-len(content) // OFFICE_CHUNK_BYTES)
        self.drop_last = drop_last
        self.calls = []

    async def command_result(self, *, command_id, run_id, owner_id, workspace_id):
        from app.claw_local_task_result_projection import LocalRunnerTerminalObservation

        if (command_id, run_id, owner_id, workspace_id) != (
                RunHistory.CORRELATION["command_id"], RunHistory.RUN,
                RunHistory.OWNER, RunHistory.WORKSPACE):
            raise RuntimeError("secret broker detail must not escape")
        return LocalRunnerTerminalObservation(
            **RunHistory.CORRELATION, run_id=RunHistory.RUN, sequence=1,
            state="acknowledged", termination="exited", exit_code=0,
            acknowledged_at="2026-10-26T10:00:00Z",
        )

    async def read_office_artifact_part(self, payload):
        self.calls.append(dict(payload))
        if (payload["owner"] != RunHistory.OWNER
                or payload["workspace"] != RunHistory.WORKSPACE
                or payload["run_id"] != RunHistory.RUN
                or payload["command_id"] != RunHistory.CORRELATION["command_id"]
                or payload["kind"] != "xlsx"):
            return _unavailable()
        index = payload["part_index"]
        if index >= self.parts or (self.drop_last and index == self.parts - 1):
            return _unavailable()
        part = self.content[index * OFFICE_CHUNK_BYTES:(index + 1) * OFFICE_CHUNK_BYTES]
        return {"ok": True, "artifact_part": {
            "contract_version": "claw-office-artifact-read.v1",
            "artifact_id": "artifact_4434_office_xlsx",
            "filename": XLSX_NAME, "media_type": XLSX_MEDIA,
            "size_bytes": len(self.content),
            "integrity_ref": sha256(self.content).hexdigest(),
            "part_index": index, "part_count": self.parts,
            "data_b64": b64encode(part).decode("ascii"),
        }}

    @property
    def _binding(self):
        # The composition reaches the Service Binding exactly this way.
        return self


def _unavailable():
    return {"ok": False, "error": {"code": "office_artifact_unavailable"}}


def composition(content, **kwargs):
    from app.claw_local_task_result_composition import LocalRunnerResultSource

    history = RunHistory(completed=kwargs.pop("completed", True))
    return LocalRunnerResultSource(
        history=history, result_port=TerminalPort(content, **kwargs)), history


def test_the_real_composition_hands_over_only_its_own_verified_workbook():
    """The routes over the real reader: run row, correlation, terminal fact, parts.

    Nothing is stubbed between a route and the private part stream, so the whole
    OOXML package is rebuilt, digest-checked and put through the same untrusted
    input classifier the PC route gate uses before a byte reaches the caller.
    """
    workbook = real_workbook(8000)
    source, history = composition(workbook)
    owned = call(OFFICE_XLSX_DOWNLOAD_PATH, source, owner=history.OWNER, run=history.RUN)
    assert owned.status_code == 200
    assert owned.body == workbook
    assert owned.headers["content-disposition"] == (
        f"attachment; filename*=UTF-8''{quote(XLSX_NAME, safe='')}"
    )
    assert source._result_port.parts > 1, "the workbook has to span chunks"
    served = [call_["part_index"] for call_ in source._result_port.calls]
    assert served == list(range(source._result_port.parts))

    other = call(OFFICE_XLSX_DOWNLOAD_PATH, source, owner=history.OTHER, run=history.RUN)
    unknown = call(OFFICE_XLSX_DOWNLOAD_PATH, source, owner=history.OWNER,
                   run="run_nobody_owns")
    assert other.status_code == unknown.status_code == 404
    assert other.body == unknown.body
    assert b"secret" not in other.body

    query = call(OFFICE_OUTPUTS_QUERY_PATH, source, owner=history.OWNER, run=history.RUN)
    assert query.status_code == 200
    listed = body(query)["outputs"]
    assert [item["media_type"] for item in listed] == [XLSX_MEDIA]
    assert listed[0]["size_bytes"] == len(workbook)
    assert listed[0]["integrity_ref"] == sha256(workbook).hexdigest()
    assert listed[0]["drive_write_approval"] is False


def test_a_run_that_never_finished_or_a_short_stream_refuses_the_download():
    workbook = real_workbook(400)
    running, history = composition(workbook, completed=False)
    assert call(OFFICE_XLSX_DOWNLOAD_PATH, running,
                owner=history.OWNER, run=history.RUN).status_code == 404

    cut, history = composition(workbook, drop_last=True)
    refused = call(OFFICE_XLSX_DOWNLOAD_PATH, cut, owner=history.OWNER, run=history.RUN)
    assert refused.status_code == 404
    assert b"PK\x03\x04" not in refused.body


def test_an_untrusted_package_does_not_become_an_owner_download():
    """Digest-consistent but not a workbook: the reader gate stops it, not the UI.

    The staged stream verifies against its own declared SHA-256, so the only thing
    between these bytes and a browser is the reader's container and classifier
    gate. A refused package must not appear in the query either.
    """
    junk = b"PK\x03\x04" + b"not a workbook" * 5000
    source, history = composition(junk)
    download = call(OFFICE_XLSX_DOWNLOAD_PATH, source, owner=history.OWNER, run=history.RUN)
    assert download.status_code in (404, 503)
    assert b"PK" not in download.body
    query = call(OFFICE_OUTPUTS_QUERY_PATH, source, owner=history.OWNER, run=history.RUN)
    assert query.status_code == 404
