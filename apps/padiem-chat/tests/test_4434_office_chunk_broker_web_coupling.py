"""#4434: publisher -> private Office staging -> server reassembly -> result.

The real KAgent chunk publisher sends real revised-workbook bytes produced by
the real Python value editor. The recorded wire is then served through the real
Chat-side private reassembly reader with the same per-part and whole-file
SHA-256 checks the deployed Durable Object store enforces, so the client wire
and the server reader are proven to agree on one contract.

A dictionary stands in for the Durable Object rows (the store has its own
SQLite-backed contract test in ``packages/padiem-control-plane``); no
production authority is weakened and no Drive WRITE is involved.
"""

from __future__ import annotations

import asyncio
import base64
from hashlib import sha256
from pathlib import Path
import sys

import pytest

from app.claw_local_office_binary_reader import (
    BrokerOfficeBinaryReader, OfficeBinaryReadRefused, OFFICE_CHUNK_BYTES,
)

from kagent.local_office_chunk_publisher import LocalOfficeChunkPublisher
from kagent.local_agent_control_plane_https import ControlPlaneHttpsOperation
from kagent.local_agent_secure_transport import (
    OutboundBrokerEndpoint, OutboundTransportConfig, OutboundTransportMode,
)
from kagent.local_xlsx_python_copy_editor import (
    IndependentZipXlsxRevisionVerifier, PythonXlsxValueCopyEditor,
)
from kagent.local_xlsx_revision_pdf import ReviewedCellRevision

import padiem_control_plane  # only to locate its sibling adapter root

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ADAPTER_ROOT = Path(padiem_control_plane.__file__).resolve().parents[1]
if str(ADAPTER_ROOT) not in sys.path:  # the store module lives beside the package
    sys.path.insert(0, str(ADAPTER_ROOT))

from local_agent_broker_office_chunks import (  # noqa: E402
    BrokerOfficeChunkStore, OfficeChunkRefused,
)

OWNER = "owner_4434"
WORKSPACE = "workspace_4434"
RUN = "run_4434_xlsx"
COMMAND = "command_4434_xlsx"
BINDING = "binding_4434_device"
ARTIFACT = "resident_xlsx_revision_result_4434"
CREDENTIAL = b"resident-device-credential"

EDIT = ReviewedCellRevision(
    sheet_name="Quote", cell="B2", previous_value=88_000_000, replacement_value=91_500_000,
)


def _revised_workbook() -> tuple[bytes, bytes]:
    from openpyxl import Workbook
    from io import BytesIO

    book = Workbook()
    sheet = book.active
    sheet.title = "Quote"
    sheet["A1"] = "기존 항목"
    sheet["B2"] = 88_000_000
    sheet["C2"] = "text kept"
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    source = buffer.getvalue()
    revised = PythonXlsxValueCopyEditor().revise_xlsx_copy(source, (EDIT,))
    assert IndependentZipXlsxRevisionVerifier().verify_only_reviewed_edits(
        source, revised, (EDIT,)) is True
    return source, revised


class RecordingTransport:
    def __init__(self) -> None:
        self.calls: list[dict] = []

    def post(self, *, config, operation, payload, timeout_seconds):
        assert config.endpoint.url.startswith("https://")
        assert operation is ControlPlaneHttpsOperation.OFFICE_PART
        self.calls.append(dict(payload))
        return {"ok": True, "office_part": {
            "stored": True, "kind": payload["kind"],
            "part_index": payload["part_index"],
        }}


def _publisher(transport: RecordingTransport) -> LocalOfficeChunkPublisher:
    return LocalOfficeChunkPublisher(
        transport=transport,
        config=OutboundTransportConfig(endpoint=OutboundBrokerEndpoint(
            endpoint_ref="broker.resident.4434", url="https://broker.example.test/",
            mode=OutboundTransportMode.HTTPS_LONG_POLL,
        )),
    )


def _publish(content: bytes) -> tuple[RecordingTransport, str]:
    transport = RecordingTransport()
    receipt = _publisher(transport).publish(
        binding_ref=BINDING, credential=CREDENTIAL,
        command_id=COMMAND, run_id=RUN, artifact_id=ARTIFACT,
        filename="book_revised.xlsx", media_type=XLSX_MIME,
        content=content, kind="xlsx",
    )
    return transport, receipt.integrity_ref


class Staging(dict):
    """The manifest plus immutable parts, keyed exactly like the store."""


def _stage(calls: list[dict]) -> Staging:
    first = calls[0]
    staging = Staging(
        artifact_id=first["artifact_id"], filename=first["filename"],
        media_type=first["media_type"], size_bytes=first["size_bytes"],
        integrity_ref=first["integrity_ref"], part_count=first["part_count"],
        command_id=first["command_id"], run_id=first["run_id"],
        binding_ref=first["binding_ref"], kind=first["kind"],
        parts={row["part_index"]: bytes(base64.b64decode(row["data_b64"]))
               for row in calls},
        part_digests={row["part_index"]: row["part_sha256"] for row in calls},
    )
    assert len(staging["parts"]) == staging["part_count"]
    return staging


class PrivateOfficeBinding:
    """Serves private part reads the way the Durable Object route does."""

    def __init__(self, staging: Staging | list[Staging]) -> None:
        self.streams = staging if isinstance(staging, list) else [staging]
        self.calls: list[dict] = []

    def _match(self, payload: dict) -> Staging | None:
        for stream in self.streams:
            if (payload["command_id"] == stream["command_id"]
                    and payload["run_id"] == stream["run_id"]
                    and payload["kind"] == stream["kind"]
                    and payload["owner"] == OWNER
                    and payload["workspace"] == WORKSPACE):
                return stream
        return None

    def read_office_artifact_part(self, payload: dict) -> dict:
        self.calls.append(dict(payload))
        stream = self._match(payload)
        if stream is None:
            return {"ok": False, "error": {"code": "office_artifact_unavailable"}}
        index = payload["part_index"]
        if index not in stream["parts"]:
            return {"ok": False, "error": {"code": "office_artifact_unavailable"}}
        return {"ok": True, "artifact_part": {
            "contract_version": "claw-office-artifact-read.v1",
            "artifact_id": stream["artifact_id"],
            "filename": stream["filename"],
            "media_type": stream["media_type"],
            "size_bytes": stream["size_bytes"],
            "integrity_ref": stream["integrity_ref"],
            "part_index": index,
            "part_count": stream["part_count"],
            "data_b64": base64.b64encode(stream["parts"][index]).decode("ascii"),
        }}


async def _read(binding: PrivateOfficeBinding, *, owner=OWNER, workspace=WORKSPACE,
                run=RUN, command=COMMAND, kind="xlsx"):
    return await BrokerOfficeBinaryReader(private_binding=binding).read(
        owner_id=owner, workspace_ref=workspace, run_ref=run, command_id=command, kind=kind,
    )


def test_published_revision_stream_matches_the_store_wire_schema() -> None:
    _source, revised = _revised_workbook()
    transport, integrity = _publish(revised)
    staging = _stage(transport.calls)

    # The store's own expectations, applied to the publisher's real payloads.
    for row in transport.calls:
        assert row["contract_version"] == "claw-office-artifact-chunk.v1"
        assert row["size_bytes"] == len(revised)
        assert row["integrity_ref"] == integrity
        assert sha256(base64.b64decode(row["data_b64"])).hexdigest() == row["part_sha256"]
        assert len(base64.b64decode(row["data_b64"])) <= OFFICE_CHUNK_BYTES
        for forbidden in ("path", "destination", "account_ref", "workspace_ref", "drive"):
            assert forbidden not in row
    assert integrity == sha256(revised).hexdigest()
    assert staging["part_count"] == len(transport.calls)


def test_server_reassembly_returns_the_exact_revised_workbook() -> None:
    source, revised = _revised_workbook()
    transport, integrity = _publish(revised)
    binding = PrivateOfficeBinding(_stage(transport.calls))

    verified = _run(_read(binding))

    assert verified.integrity_ref == integrity
    assert verified.content == revised
    assert verified.artifact_id == ARTIFACT
    assert verified.filename == "book_revised.xlsx"
    # The original bytes are not the result, and the result opens as a workbook.
    assert verified.content != source
    from io import BytesIO
    from openpyxl import load_workbook
    sheet = load_workbook(BytesIO(verified.content))["Quote"]
    assert sheet["B2"].value == 91_500_000
    assert sheet["A1"].value == "기존 항목"
    assert sheet["C2"].value == "text kept"


def test_two_streams_for_two_commands_are_read_back_separately() -> None:
    _source, revised = _revised_workbook()
    first_transport, _first = _publish(revised)
    streams = [_stage(first_transport.calls)]
    second_content = revised
    other = "command_4434_other"
    transport2 = RecordingTransport()
    _publisher(transport2).publish(
        binding_ref=BINDING, credential=CREDENTIAL,
        command_id=other, run_id="run_4434_other", artifact_id="artifact_4434_other",
        filename="second_revised.xlsx", media_type=XLSX_MIME,
        content=second_content, kind="xlsx",
    )
    streams.append(_stage(transport2.calls))
    binding = PrivateOfficeBinding(streams)

    one = _run(_read(binding))
    two = _run(_read(binding, run="run_4434_other", command=other))

    assert one.content == revised and two.content == second_content
    assert one.artifact_id != two.artifact_id
    assert two.filename == "second_revised.xlsx"


def test_foreign_owner_run_command_and_kind_are_refused() -> None:
    _source, revised = _revised_workbook()
    transport, _integrity = _publish(revised)
    binding = PrivateOfficeBinding(_stage(transport.calls))

    with pytest.raises(OfficeBinaryReadRefused):
        _run(_read(binding, owner="owner_other"))
    with pytest.raises(OfficeBinaryReadRefused):
        _run(_read(binding, workspace="workspace_other"))
    with pytest.raises(OfficeBinaryReadRefused):
        _run(_read(binding, run="run_other", command=COMMAND))
    with pytest.raises(OfficeBinaryReadRefused):
        _run(_read(binding, kind="pdf"))


def test_tampered_part_and_whole_file_digests_are_refused() -> None:
    _source, revised = _revised_workbook()
    transport, _integrity = _publish(revised)
    staging = _stage(transport.calls)

    tampered = dict(staging)
    parts = dict(staging["parts"])
    parts[0] = parts[0][:-4] + b"XXXX"
    tampered["parts"] = parts
    with pytest.raises(OfficeBinaryReadRefused):
        _run(_read(PrivateOfficeBinding(Staging(tampered))))

    shifted = dict(staging)
    shifted["integrity_ref"] = sha256(revised + b"tail").hexdigest()
    with pytest.raises(OfficeBinaryReadRefused):
        _run(_read(PrivateOfficeBinding(Staging(shifted))))


def test_missing_part_stops_the_read_and_a_resumed_stream_recovers() -> None:
    _source, revised = _revised_workbook()
    transport, _integrity = _publish(revised)
    staging = _stage(transport.calls)
    assert staging["part_count"] >= 1

    truncated = Staging(dict(staging))
    truncated["parts"] = {index: part for index, part in staging["parts"].items()
                          if index != staging["part_count"] - 1}
    truncated["part_count"] = staging["part_count"]
    with pytest.raises(OfficeBinaryReadRefused):
        _run(_read(PrivateOfficeBinding(truncated)))

    # A resumed delivery of the same stream recovers without a second artifact.
    recovered = _run(_read(PrivateOfficeBinding(staging)))
    assert recovered.content == revised
    assert recovered.artifact_id == ARTIFACT


def test_store_module_is_the_same_contract_the_publisher_uses() -> None:
    """Pin the shared wire fields the Durable Object store validates."""
    from local_agent_broker_office_chunks import _WIRE as STORE_WIRE

    _source, revised = _revised_workbook()
    transport, _integrity = _publish(revised)
    # The device route keeps ownership out of the payload: it removes the
    # binding and credential and passes only artifact material to the store.
    material = {
        key: value for key, value in transport.calls[0].items()
        if key not in ("binding_ref", "credential_b64")
    }
    assert frozenset(material) == STORE_WIRE
    # Building the real store against a storage without .sql must refuse, which
    # is the durable-storage gate the deployed Durable Object satisfies.
    with pytest.raises(OfficeChunkRefused):
        BrokerOfficeChunkStore(storage=object(), state_port=None,
                               authority_ref="authority_4434")


def _run(coroutine):
    return asyncio.run(coroutine)
