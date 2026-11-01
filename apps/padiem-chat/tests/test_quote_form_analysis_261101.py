"""Three synthetic layouts exercise suggestions from actual XLSX structure."""
from io import BytesIO
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "padiem-file-lane" / "src"))
from padiem_file_lane.xlsx_edit import describe
from app.b66_quote_form_analysis import analyze_quote_form


@pytest.mark.parametrize("shift,head,labels", [
    (0, 9, ("품명", "수량", "단가", "금액", "소계", "부가세", "총계", "견적일", "수신처")),
    (1, 12, ("item", "qty", "unit price", "amount", "subtotal", "vat", "total", "date", "customer")),
    (2, 15, ("품목명", "수량", "단가", "금액", "공급가액", "부가가치세", "합계", "견적일자", "거래처")),
])
def test_unrelated_layouts_from_xlsx(shift, head, labels):
    book = Workbook()
    sheet = book.active
    sheet.title = "Fictional"
    sheet.cell(2, 6, "Fictional logo")
    for offset, label in enumerate(labels[:4]):
        sheet.cell(head, shift + offset + 1, label)
    for offset, label in enumerate(labels[4:7]):
        sheet.cell(head + 4 + offset, shift + 3, label)
        sheet.cell(head + 4 + offset, shift + 4, "=1+1")
    sheet.cell(5, shift + 1, labels[7])
    sheet.cell(7, shift + 1, labels[8])
    output = BytesIO()
    book.save(output)
    actual = analyze_quote_form(describe(output.getvalue()))
    def cell(column, row):
        from openpyxl.utils import get_column_letter
        return get_column_letter(column) + str(row)
    expected = {
        "item": cell(shift + 1, head + 1),
        "quantity": cell(shift + 2, head + 1),
        "unit_price": cell(shift + 3, head + 1),
        "amount": cell(shift + 4, head + 1),
        "subtotal": cell(shift + 4, head + 4),
        "tax": cell(shift + 4, head + 5),
        "total": cell(shift + 4, head + 6),
        "date": cell(shift + 2, 5),
        "recipient": cell(shift + 2, 7),
    }
    for role, target in expected.items():
        assert not actual[role]["requires_confirmation"], (role, actual[role])
        assert actual[role]["candidates"][0]["cell"] == target


def test_duplicate_labels_are_not_silently_resolved():
    structure = {"sheets": [{"name": "Fictional", "merged": [], "cells": [
        {"cell": "A4", "text": "단가", "formula": False},
        {"cell": "D4", "text": "단가", "formula": False},
    ]}]}
    unit = analyze_quote_form(structure)["unit_price"]
    assert unit["requires_confirmation"] is True
    assert {item["cell"] for item in unit["candidates"]} == {"A5", "D5"}
    assert analyze_quote_form(structure)["total"]["requires_confirmation"] is True


def test_existing_text_and_merged_label_not_overwritten():
    structure = {"sheets": [{"name": "Fictional", "merged": ["A1:B1"], "cells": [
        {"cell": "A1", "text": "견적일", "formula": False},
        {"cell": "D2", "text": "수신처", "formula": False},
        {"cell": "E2", "text": "Existing text", "formula": False},
    ]}]}
    result = analyze_quote_form(structure)
    assert result["date"]["requires_confirmation"] and not result["date"]["candidates"]
    assert result["recipient"]["requires_confirmation"] and not result["recipient"]["candidates"]
