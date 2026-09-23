from __future__ import annotations

import inspect
import os
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest

try:
    from openpyxl import Workbook
except ModuleNotFoundError:  # Core base install intentionally excludes document extras.
    Workbook = None

try:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
except ModuleNotFoundError:  # Core base install intentionally excludes document extras.
    PdfWriter = None
    DecodedStreamObject = DictionaryObject = NameObject = None

from padiem_ai_core.document_normalization import (
    MAX_BINARY_DOCUMENT_BYTES,
    MAX_DOCUMENT_CHARS,
    MAX_OOXML_ENTRIES,
    MAX_OOXML_ENTRY_UNCOMPRESSED_BYTES,
    MAX_OOXML_TOTAL_UNCOMPRESSED_BYTES,
    MAX_PDF_PAGES,
    BINARY_DOCUMENT_MEDIA,
    DocumentNormalizationError,
    _safe_ooxml_member,
    extract_binary_document,
    extract_docx_text,
    extract_hwpx_text,
    extract_pptx_text,
    normalize_text_document,
    parse_hwpx_sections,
    validate_document_identity,
    validate_ooxml_archive,
)
import padiem_ai_core.document_normalization as document_normalization
from padiem_ai_core.hwpx_package_serializer import (
    HwpxPackageSection,
    deserialize_hwpx_package,
)

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HWPX_MIME = "application/hwp+zip"


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    """Build an in-memory ZIP whose member names survive verbatim.

    ``ZipInfo.__init__`` rewrites ``os.sep`` to ``/``, so on Windows a member
    named ``a\\b.xml`` was silently stored as ``a/b.xml``. The unsafe-path guard
    that rejects a backslash member was therefore never actually exercised on
    Windows, and the adversarial fixture was quietly not adversarial there.
    Assigning the name *after* construction keeps the raw member name on every
    platform, which is what the guard has to be tested against.
    """

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            member = ZipInfo(name)
            member.filename = name
            # ``writestr(str, ...)`` sets this from the archive default; passing
            # a ZipInfo does not, so it is set here to keep the compression
            # behaviour of every existing fixture exactly as it was.
            member.compress_type = ZIP_DEFLATED
            archive.writestr(member, payload)
    return output.getvalue()


def _mark_first_entry_encrypted(payload: bytes) -> bytes:
    data = bytearray(payload)
    for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        index = data.find(signature)
        assert index >= 0
        start = index + flag_offset
        flags = int.from_bytes(data[start : start + 2], "little") | 0x1
        data[start : start + 2] = flags.to_bytes(2, "little")
    return bytes(data)


def _docx_xml(*paragraphs: str) -> bytes:
    body = "".join(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" for text in paragraphs)
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    ).encode()


def _pptx_xml(*parts: str) -> bytes:
    body = "".join(f"<a:r><a:t>{text}</a:t></a:r>" for text in parts)
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<p:cSld><p:spTree><p:sp><p:txBody><a:p>{body}</a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    ).encode()


def _hwpx_section_xml(*paragraphs: str) -> bytes:
    body = "".join(f"<hp:p><hp:runs><hp:t>{text}</hp:t></hp:runs></hp:p>" for text in paragraphs)
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
        'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
        f"{body}</hs:sec>"
    ).encode()


def _require_pdf_extra() -> None:
    if PdfWriter is None:
        pytest.skip("pypdf is provided by the optional documents extra")


def _require_xlsx_extra() -> None:
    if Workbook is None:
        pytest.skip("openpyxl is provided by the optional documents extra")


def _minimal_text_pdf(text: str) -> bytes:
    _require_pdf_extra()
    assert PdfWriter is not None
    assert DictionaryObject is not None and NameObject is not None and DecodedStreamObject is not None
    writer = PdfWriter()
    page = writer.add_blank_page(width=320, height=180)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    content = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content.set_data(f"BT /F1 14 Tf 36 90 Td ({escaped}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _blank_pdf(*, encrypted: bool = False, pages: int = 1) -> bytes:
    _require_pdf_extra()
    assert PdfWriter is not None
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=320, height=180)
    if encrypted:
        writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    _require_xlsx_extra()
    assert Workbook is not None
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "Padiem XLSX"
    sheet["B2"] = 42
    detail = workbook.create_sheet("Detail")
    detail["C3"] = "bounded"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _replace_zip_entry(payload: bytes, target: str, replacement: bytes) -> bytes:
    with ZipFile(BytesIO(payload)) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries[target] = replacement
    return _zip_bytes(entries)


def test_text_normalization_is_bounded_and_public_projection_is_safe() -> None:
    document = normalize_text_document(
        name=" notes.txt ",
        media_type="text/plain",
        text="\ufefffirst\r\nsecond\rthird",
    )
    assert document.name == "notes.txt"
    assert document.text == "first\nsecond\nthird"
    assert document.source_kind == "text"
    assert document.to_public_dict() == {
        "type": "document",
        "name": "notes.txt",
        "media_type": "text/plain",
        "byte_size": len(document.text.encode("utf-8")),
        "text_chars": len(document.text),
    }
    assert document.text not in repr(document)


@pytest.mark.parametrize(
    ("name", "media_type", "code"),
    [
        ("notes.pdf", "text/plain", "media_extension_mismatch"),
        ("notes.txt", "application/pdf", "unsupported_text_media_type"),
        ("legacy.hwp", "application/x-hwp", "unsupported_text_media_type"),
        ("legacy.hwpx", "application/zip", "unsupported_text_media_type"),
    ],
)
def test_text_identity_fails_closed(name: str, media_type: str, code: str) -> None:
    with pytest.raises(DocumentNormalizationError) as exc:
        normalize_text_document(name=name, media_type=media_type, text="safe")
    assert exc.value.code == code


def test_text_rejects_nul_controls_and_character_overflow_without_echo() -> None:
    secret = "PRIVATE-DOCUMENT-SECRET"
    with pytest.raises(DocumentNormalizationError) as nul:
        normalize_text_document(name="x.txt", media_type="text/plain", text=secret + "\x00")
    assert nul.value.code == "binary_text_rejected"
    assert secret not in str(nul.value)

    with pytest.raises(DocumentNormalizationError) as controls:
        normalize_text_document(name="x.txt", media_type="text/plain", text="safe\x01\x02\x03\x04")
    assert controls.value.code == "excessive_control_characters"

    with pytest.raises(DocumentNormalizationError) as large:
        normalize_text_document(name="x.txt", media_type="text/plain", text="x" * (MAX_DOCUMENT_CHARS + 1))
    assert large.value.code == "text_too_long"


def test_binary_identity_is_product_neutral_and_has_no_path_or_url_authority() -> None:
    name, media = validate_document_identity(name="sample.pdf", media_type=PDF_MIME, source_kind="binary")
    assert (name, media) == ("sample.pdf", PDF_MIME)
    assert set(BINARY_DOCUMENT_MEDIA) == {PDF_MIME, DOCX_MIME, PPTX_MIME, XLSX_MIME, HWPX_MIME}
    with pytest.raises(DocumentNormalizationError) as mismatch:
        validate_document_identity(name="sample.docx", media_type=PDF_MIME, source_kind="binary")
    assert mismatch.value.code == "media_extension_mismatch"

    import inspect

    assert tuple(inspect.signature(extract_binary_document).parameters) == ("name", "media_type", "payload")


def test_pdf_extracts_text_and_never_retains_raw_binary() -> None:
    payload = _minimal_text_pdf("Padiem Core PDF")
    document = extract_binary_document(name="sample.pdf", media_type=PDF_MIME, payload=payload)
    assert "Padiem Core PDF" in document.text
    assert document.byte_size == len(payload)
    assert document.source_kind == "binary"
    assert repr(payload[:24]) not in repr(document)
    assert "Padiem Core PDF" not in repr(document)
    assert "text" not in document.to_public_dict()


def test_pdf_bad_magic_encryption_page_limit_and_scanned_only_fail_closed() -> None:
    with pytest.raises(DocumentNormalizationError) as magic:
        extract_binary_document(name="x.pdf", media_type=PDF_MIME, payload=b"not a pdf")
    assert magic.value.code == "pdf_magic_mismatch"

    _require_pdf_extra()
    with pytest.raises(DocumentNormalizationError) as encrypted:
        extract_binary_document(name="x.pdf", media_type=PDF_MIME, payload=_blank_pdf(encrypted=True))
    assert encrypted.value.code == "pdf_encrypted"

    with pytest.raises(DocumentNormalizationError) as pages:
        extract_binary_document(name="x.pdf", media_type=PDF_MIME, payload=_blank_pdf(pages=MAX_PDF_PAGES + 1))
    assert pages.value.code == "pdf_page_limit"

    with pytest.raises(DocumentNormalizationError) as scanned:
        extract_binary_document(name="x.pdf", media_type=PDF_MIME, payload=_blank_pdf())
    assert scanned.value.code == "pdf_empty_text"


def test_docx_preserves_paragraphs_and_pptx_preserves_slide_order() -> None:
    docx = _zip_bytes({"word/document.xml": _docx_xml("paragraph one", "paragraph two")})
    document = extract_binary_document(name="sample.docx", media_type=DOCX_MIME, payload=docx)
    assert document.text == "paragraph one\nparagraph two"

    pptx = _zip_bytes(
        {
            "ppt/slides/slide2.xml": _pptx_xml("slide two", "second line"),
            "ppt/slides/slide1.xml": _pptx_xml("slide one"),
        }
    )
    presentation = extract_binary_document(name="slides.pptx", media_type=PPTX_MIME, payload=pptx)
    assert presentation.text == "slide one\nslide two\nsecond line"


def test_hwpx_round_trip_orders_sections_numerically() -> None:
    payload = _zip_bytes(
        {
            "mimetype": b"application/hwp+zip",
            "Contents/section2.xml": _hwpx_section_xml("second section", "second body"),
            "Contents/section10.xml": _hwpx_section_xml("tenth section"),
            "Contents/section1.xml": _hwpx_section_xml("first section"),
        }
    )
    document = extract_binary_document(name="report.hwpx", media_type=HWPX_MIME, payload=payload)
    assert document.text == "first section\nsecond section\nsecond body\ntenth section"
    assert document.source_kind == "binary"
    assert document.byte_size == len(payload)
    assert "text" not in document.to_public_dict()


def test_hwpx_captures_table_cell_paragraphs_without_duplication() -> None:
    section = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
        'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
        "<hp:p><hp:runs><hp:t>intro</hp:t></hp:runs></hp:p>"
        "<hp:p><hp:tbl><hp:tr>"
        "<hp:tc><hp:p><hp:runs><hp:t>cell-a</hp:t></hp:runs></hp:p></hp:tc>"
        "<hp:tc><hp:p><hp:runs><hp:t>cell-b</hp:t></hp:runs></hp:p></hp:tc>"
        "</hp:tr></hp:tbl></hp:p>"
        "</hs:sec>"
    ).encode()
    payload = _zip_bytes({"Contents/section1.xml": section})
    document = extract_binary_document(name="table.hwpx", media_type=HWPX_MIME, payload=payload)
    assert document.text == "intro\ncell-a\ncell-b"


def test_hwpx_missing_part_empty_and_mimetype_fail_closed() -> None:
    with pytest.raises(DocumentNormalizationError) as missing:
        extract_hwpx_text(_zip_bytes({"mimetype": b"application/hwp+zip"}))
    assert missing.value.code == "hwpx_missing_part"

    with pytest.raises(DocumentNormalizationError) as empty:
        extract_hwpx_text(_zip_bytes({"Contents/section1.xml": _hwpx_section_xml()}))
    assert empty.value.code == "hwpx_empty"

    with pytest.raises(DocumentNormalizationError) as mismatch:
        extract_hwpx_text(
            _zip_bytes(
                {
                    "mimetype": b"application/zip",
                    "Contents/section1.xml": _hwpx_section_xml("safe"),
                }
            )
        )
    assert mismatch.value.code == "hwpx_mimetype_mismatch"

    without_mimetype = _zip_bytes({"Contents/section1.xml": _hwpx_section_xml("safe")})
    assert extract_hwpx_text(without_mimetype) == "safe"


def test_hwpx_encrypted_archive_fails_closed_and_legacy_hwp_stays_unsupported() -> None:
    encrypted = _mark_first_entry_encrypted(
        _zip_bytes(
            {
                "mimetype": b"application/hwp+zip",
                "Contents/section1.xml": _hwpx_section_xml("safe"),
            }
        )
    )
    with pytest.raises(DocumentNormalizationError) as locked:
        extract_hwpx_text(encrypted)
    assert locked.value.code == "ooxml_encrypted"

    with pytest.raises(DocumentNormalizationError) as legacy:
        extract_binary_document(name="legacy.hwp", media_type="application/x-hwp", payload=b"\xd0\xcf\x11\xe0fake")
    assert legacy.value.code == "unsupported_binary_media_type"

    with pytest.raises(DocumentNormalizationError) as extension:
        extract_binary_document(name="notes.txt", media_type=HWPX_MIME, payload=b"PK\x03\x04")
    assert extension.value.code == "media_extension_mismatch"


# --------------------------------------------------------------------------- #
# #2966: parse_hwpx_sections is now the single HWPX archive-and-XML read
# authority. These cases pin the facts it reports, the flat projection it feeds,
# and the fact that the flat path no longer walks an archive of its own.
# --------------------------------------------------------------------------- #

_HWPX_NS = (
    'xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
    'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'
)


def _hwpx_raw_section(body: str) -> bytes:
    return f'<?xml version="1.0" encoding="UTF-8"?><hs:sec {_HWPX_NS}>{body}</hs:sec>'.encode()


def _hwpx_para(text: str) -> str:
    return f"<hp:p><hp:runs><hp:t>{text}</hp:t></hp:runs></hp:p>"


def _hwpx_payload(*section_bodies: bytes) -> bytes:
    entries = {"mimetype": b"application/hwp+zip"}
    entries.update(
        {f"Contents/section{index}.xml": body for index, body in enumerate(section_bodies, start=1)}
    )
    return _zip_bytes(entries)


def test_parse_hwpx_sections_reports_shape_the_flat_text_flattens() -> None:
    table = (
        _hwpx_para("intro")
        + "<hp:p><hp:tbl><hp:tr><hp:tc>"
        + _hwpx_para("cell-a")
        + "</hp:tc></hp:tr></hp:tbl></hp:p>"
    )
    payload = _hwpx_payload(_hwpx_raw_section(table))
    (section,) = parse_hwpx_sections(payload)
    assert [(item.text, item.text_nodes, item.holds_nested_paragraph) for item in section.paragraphs] == [
        ("intro", 1, False),
        ("cell-a", 1, True),
        ("cell-a", 1, False),
    ]
    assert section.root_tag.endswith("}sec")
    assert section.unsupported_nodes == 3
    assert section.has_carriage_return is False


def test_parse_hwpx_sections_orders_by_section_number_not_member_name() -> None:
    sequential = _hwpx_payload(
        _hwpx_section_xml("one"),
        _hwpx_section_xml("two"),
        _hwpx_section_xml("three"),
    )
    entries = {
        "mimetype": b"application/hwp+zip",
        "Contents/section1.xml": _hwpx_section_xml("one"),
        "Contents/section10.xml": _hwpx_section_xml("ten"),
        "Contents/section2.xml": _hwpx_section_xml("two"),
    }
    assert [section.index for section in parse_hwpx_sections(_zip_bytes(entries))] == [1, 2, 10]
    assert [section.index for section in parse_hwpx_sections(sequential)] == [1, 2, 3]


def test_parse_hwpx_sections_flags_a_carriage_return_the_parser_would_erase() -> None:
    with_return = _hwpx_payload(_hwpx_raw_section(_hwpx_para("일\r\n이")))
    without_return = _hwpx_payload(_hwpx_raw_section(_hwpx_para("일\n이")))
    assert parse_hwpx_sections(with_return)[0].has_carriage_return is True
    assert parse_hwpx_sections(without_return)[0].has_carriage_return is False
    # Both still flatten: the flat authority never promised byte-exact control
    # characters, and its behaviour must not move.
    assert extract_hwpx_text(with_return) == extract_hwpx_text(without_return) == "일\n이"


def test_flat_hwpx_text_projection_is_pinned_after_the_factoring() -> None:
    cases: list[tuple[bytes, str]] = [
        (_hwpx_payload(_hwpx_section_xml("alone")), "alone"),
        (_hwpx_payload(_hwpx_section_xml("a", "b")), "a\nb"),
        (_hwpx_payload(_hwpx_section_xml("", "b")), "b"),
        (_hwpx_payload(_hwpx_section_xml("  a  ", "  b  ")), "a  \n  b"),
        (_hwpx_payload(_hwpx_section_xml(), _hwpx_section_xml("본문")), "본문"),
        (_hwpx_payload(_hwpx_raw_section("<hp:p><hp:runs><hp:t>앞</hp:t></hp:runs>"
                                         "<hp:runs><hp:t>뒤</hp:t></hp:runs></hp:p>")), "앞뒤"),
        (_hwpx_payload(_hwpx_raw_section("<hp:p/>" + _hwpx_para("본문"))), "본문"),
        (_hwpx_payload(_hwpx_raw_section('<hp:p><hp:runs><hp:img href="rId1"/></hp:runs></hp:p>'
                                         + _hwpx_para("본문"))), "본문"),
        (_hwpx_payload(_hwpx_raw_section("<hp:p><hp:sub>" + _hwpx_para("a") + _hwpx_para("b")
                                         + "</hp:sub>tail</hp:p>")), "a\nb"),
        (_hwpx_payload(_hwpx_raw_section(_hwpx_para("일\r\n이") + _hwpx_para("둘"))), "일\n이" + "\n" + "둘"),
    ]
    for payload, expected in cases:
        assert extract_hwpx_text(payload) == expected

    # A paragraph that carries text directly instead of through hp:t has always
    # produced nothing, and the factoring keeps it that way.
    direct_text = _hwpx_payload(
        _hwpx_raw_section("<hp:p><hp:sub><hp:p>a</hp:p><hp:p>b</hp:p></hp:sub>tail</hp:p>")
    )
    with pytest.raises(DocumentNormalizationError) as unreadable:
        extract_hwpx_text(direct_text)
    assert unreadable.value.code == "hwpx_empty"


def test_extract_hwpx_text_keeps_no_archive_or_xml_walk_of_its_own() -> None:
    source = inspect.getsource(document_normalization)
    flat_path = source.split("def extract_hwpx_text")[1].split("\ndef ")[0]
    for forbidden in ("ZipFile(", "namelist(", "_parse_xml(", "validate_ooxml_archive(", "archive.read("):
        assert forbidden not in flat_path, forbidden
    assert "parse_hwpx_sections(" in flat_path


def test_both_hwpx_projections_route_through_the_same_single_parse_path(monkeypatch: pytest.MonkeyPatch) -> None:
    import padiem_ai_core.hwpx_package_serializer as serializer_module

    payload = _hwpx_payload(_hwpx_section_xml("단일 경로"), _hwpx_section_xml("둘째"))
    calls: list[bytes] = []
    real = document_normalization.parse_hwpx_sections

    def counting(parsed_payload: bytes):
        calls.append(parsed_payload)
        return real(parsed_payload)

    monkeypatch.setattr(document_normalization, "parse_hwpx_sections", counting)
    monkeypatch.setattr(serializer_module, "parse_hwpx_sections", counting)

    assert document_normalization.extract_hwpx_text(payload) == "단일 경로" + "\n" + "둘째"
    assert len(calls) == 1
    assert deserialize_hwpx_package(payload).sections == (
        HwpxPackageSection(paragraphs=("단일 경로",)),
        HwpxPackageSection(paragraphs=("둘째",)),
    )
    assert len(calls) == 2
    assert calls == [payload, payload]


def test_ooxml_malformed_missing_dtd_encryption_and_paths_fail_closed() -> None:
    with pytest.raises(DocumentNormalizationError) as malformed:
        extract_docx_text(b"not-a-zip")
    assert malformed.value.code == "ooxml_malformed"

    with pytest.raises(DocumentNormalizationError) as missing:
        extract_docx_text(_zip_bytes({"safe.xml": b"<x/>"}))
    assert missing.value.code == "docx_missing_part"

    dtd = _zip_bytes({"word/document.xml": b'<!DOCTYPE x [<!ENTITY e "private">]><x>&e;</x>'})
    with pytest.raises(DocumentNormalizationError) as rejected:
        extract_docx_text(dtd)
    assert rejected.value.code == "ooxml_dtd_rejected"
    assert "private" not in str(rejected.value)

    encrypted = _mark_first_entry_encrypted(_zip_bytes({"word/document.xml": _docx_xml("safe")}))
    with pytest.raises(DocumentNormalizationError) as locked:
        extract_docx_text(encrypted)
    assert locked.value.code == "ooxml_encrypted"

    for unsafe in ("../evil.xml", "/absolute.xml", "C:/drive.xml", "a//b.xml"):
        payload = _zip_bytes({unsafe: b"x", "word/document.xml": _docx_xml("safe")})
        with pytest.raises(DocumentNormalizationError) as path_error:
            extract_docx_text(payload)
        assert path_error.value.code == "ooxml_unsafe_path"


def test_ooxml_backslash_member_guard_is_actually_executed() -> None:
    """The backslash branch of the unsafe-path guard must really run.

    ``zipfile`` normalizes ``os.sep`` to ``/`` on **read** as well as write:
    ``ZipFile._RealGetContents`` builds ``ZipInfo(name)`` from the on-disk name,
    so on Windows a backslash member can never reach the guard through an archive
    at all. Previously this case sat inside the archive loop above, where it
    silently stopped being adversarial on Windows: the document was admitted
    because the fixture had been rewritten, not because the guard was bypassed.

    So the guard is asserted directly — which executes it on every platform —
    and the archive-level behaviour is asserted per platform instead of
    disappearing from coverage. No Core parser source changes for this.
    """

    assert _safe_ooxml_member("a" + chr(92) + "b.xml") is False
    # Positive control: the assertion above is about the backslash, not about
    # the helper rejecting everything.
    assert _safe_ooxml_member("a/b.xml") is True

    archive = _zip_bytes({"a\\b.xml": b"x", "word/document.xml": _docx_xml("safe")})
    with ZipFile(BytesIO(archive)) as opened:
        stored = opened.namelist()[0]

    if os.sep == "\\":
        # The member is written with a real backslash, and the stdlib reader is
        # what neutralises it before the guard can see it.
        assert stored == "a/b.xml"
        assert extract_docx_text(archive) == "safe"
    else:
        assert stored == "a\\b.xml"
        with pytest.raises(DocumentNormalizationError) as path_error:
            extract_docx_text(archive)
        assert path_error.value.code == "ooxml_unsafe_path"


def test_ooxml_entry_count_and_uncompressed_bounds_fail_closed() -> None:
    entries = {f"safe/{index}.xml": b"x" for index in range(MAX_OOXML_ENTRIES + 1)}
    with pytest.raises(DocumentNormalizationError) as count:
        validate_ooxml_archive(_zip_bytes(entries))
    assert count.value.code == "ooxml_entry_count"

    oversized = _zip_bytes({"oversized.bin": b"x" * (MAX_OOXML_ENTRY_UNCOMPRESSED_BYTES + 1)})
    with pytest.raises(DocumentNormalizationError) as entry:
        validate_ooxml_archive(oversized)
    assert entry.value.code == "ooxml_entry_size"

    chunk = b"x" * MAX_OOXML_ENTRY_UNCOMPRESSED_BYTES
    total_entries = {
        f"safe/{index}.bin": chunk
        for index in range((MAX_OOXML_TOTAL_UNCOMPRESSED_BYTES // MAX_OOXML_ENTRY_UNCOMPRESSED_BYTES) + 1)
    }
    with pytest.raises(DocumentNormalizationError) as total:
        validate_ooxml_archive(_zip_bytes(total_entries))
    assert total.value.code == "ooxml_total_size"


def test_xlsx_extracts_values_and_archive_dtd_is_rejected_before_openpyxl() -> None:
    payload = _xlsx_bytes()
    document = extract_binary_document(name="book.xlsx", media_type=XLSX_MIME, payload=payload)
    assert "[Summary!A1] Padiem XLSX" in document.text
    assert "[Summary!B2] 42" in document.text
    assert "[Detail!C3] bounded" in document.text

    malicious = _replace_zip_entry(
        payload,
        "xl/workbook.xml",
        b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e "private">]><workbook/>',
    )
    with pytest.raises(DocumentNormalizationError) as dtd:
        extract_binary_document(name="book.xlsx", media_type=XLSX_MIME, payload=malicious)
    assert dtd.value.code == "ooxml_dtd_rejected"
    assert "private" not in str(dtd.value)


def test_binary_size_bound_fails_without_echoing_payload() -> None:
    payload = b"PRIVATE" + b"x" * MAX_BINARY_DOCUMENT_BYTES
    with pytest.raises(DocumentNormalizationError) as exc:
        extract_binary_document(name="x.pdf", media_type=PDF_MIME, payload=payload)
    assert exc.value.code == "binary_too_large"
    assert "PRIVATE" not in str(exc.value)
